"""
Spreadsheets as first-class models (§29.9): lift, check against the workbook, refuse the rest.

Workbooks are built with openpyxl at test time. openpyxl never calculates, so
where a test needs the workbook's own cached results (what Excel stores on each
recalculation) ``cached()`` writes them into the file the way Excel does.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from typing import Any

import numpy as np
import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

from maya.core.errors import ValidationFailed  # noqa: E402
from maya.formula.codegen import compile_reference  # noqa: E402
from maya.formula.evaluate import evaluate, ncdf, npdf  # noqa: E402
from maya.formula.ir import validate_ir  # noqa: E402
from maya.formula.xlsx import lift_workbook, tokenize  # noqa: E402


# -- building workbooks ----------------------------------------------------------------------
def book(
    cells: dict[str, Any],
    sheets: dict[str, dict[str, Any]] | None = None,
    names: dict[str, str] | None = None,
) -> bytes:
    """cells go on sheet 'Model'; names map name -> 'Sheet!$B$2' targets."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    for coord, value in cells.items():
        ws[coord] = value
    for title, extra in (sheets or {}).items():
        other = wb.create_sheet(title)
        for coord, value in extra.items():
            other[coord] = value
    for name, target in (names or {}).items():
        wb.defined_names[name] = DefinedName(name, attr_text=target)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cached(data: bytes, values: dict[str, Any], sheet_no: int = 1) -> bytes:
    """Write Excel's cached results (<v>) into formula cells, as a recalculation would."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == f"xl/worksheets/sheet{sheet_no}.xml":
                xml = body.decode()
                for coord, v in values.items():
                    kind = ' t="e"' if isinstance(v, str) else ""
                    xml, n = re.subn(
                        rf'<c r="{coord}"([^>]*)><f>(.*?)</f><v\s*>\s*</v></c>',
                        rf'<c r="{coord}"\1{kind}><f>\2</f><v>{v}</v></c>',
                        xml,
                    )
                    assert n == 1, coord
                body = xml.encode()
            dst.writestr(item, body)
    return out.getvalue()


LOAN = {
    "A1": "principal",
    "B1": 250000,
    "A2": "annual rate",
    "B2": 0.06,
    "A3": "years",
    "B3": 30,
    "B4": "=B2/12",
    "B5": "=B3*12",
    "B6": "=B1*B4/(1-(1+B4)^-B5)",
    "B7": "=IF(B6>1400, B6*VLOOKUP(B3, Fees!A1:B3, 2), 0)",
}
FEES = {"A1": 0, "B1": 1.0, "A2": 15, "B2": 1.01, "A3": 25, "B3": 1.02}
LOAN_NAMES = {"principal": "Model!$B$1", "rate": "Model!$B$2"}


def loan_values() -> dict[str, float]:
    r, n = 0.06 / 12, 360
    pay = 250000 * r / (1 - (1 + r) ** -n)
    return {"B4": r, "B5": n, "B6": pay, "B7": pay * 1.02}


def one(formula: str, cells: dict[str, Any] | None = None, **kw: Any) -> dict[str, Any]:
    return lift_workbook(book({**(cells or {}), "Z1": formula}), output="Z1", **kw)


def value_of(ir: dict[str, Any], params: dict[str, Any] | None = None) -> float:
    wb = ir["lifted_from"]["workbook"]
    inputs = {k: np.array([v]) for k, v in wb["values"].items()}
    out = evaluate(ir, inputs, params or {})[ir["outputs"][0]["name"]]
    return float(out[0])


# -- the lift ---------------------------------------------------------------------------------
def test_a_loan_workbook_lifts_with_names_roles_and_a_lookup():
    ir = lift_workbook(book(LOAN, {"Fees": FEES}, LOAN_NAMES), roles={"rate": "parameter"})
    assert validate_ir(ir) == []
    assert ir["outputs"] == [{"name": "B7", "type": "float64"}]
    assert {i["name"]: i["role"] for i in ir["inputs"]} == {
        "B3": "feature",
        "principal": "feature",
        "rate": "parameter",
    }
    assert set(ir["lets"]) == {"B4", "B5", "B6"}
    wb = ir["lifted_from"]["workbook"]
    assert wb["output"] == "Model!B7" and wb["cells"]["rate"] == "Model!B2"
    assert wb["values"] == {"B3": 30.0, "principal": 250000.0, "rate": 0.06}
    assert wb["lookups"] == [
        {"cell": "Model!B7", "function": "VLOOKUP", "rows": 3, "match": "approximate"}
    ]
    assert wb["check"]["status"] == "unchecked" and "never recalculated" in wb["check"]["statement"]
    assert len(wb["sha256"]) == 64 and r"\mathit{rate}" in ir["latex"]
    assert value_of(ir, {"rate": 0.06}) == pytest.approx(loan_values()["B7"], rel=1e-12)


def test_the_lift_agrees_with_the_workbooks_own_cached_results():
    data = cached(book(LOAN, {"Fees": FEES}, LOAN_NAMES), loan_values())
    check = lift_workbook(data)["lifted_from"]["workbook"]["check"]
    assert check["status"] == "agreed" and check["checked"] == check["formula_cells"] == 4
    assert "4 of 4" in check["statement"]


def test_a_workbook_that_disagrees_is_reported_cell_by_cell():
    wrong = {**loan_values(), "B6": loan_values()["B6"] + 1}
    check = lift_workbook(cached(book(LOAN, {"Fees": FEES}), wrong))["lifted_from"]["workbook"][
        "check"
    ]
    assert check["status"] == "disagreed" and "blocked" in check["statement"]
    assert [d["cell"] for d in check["disagreements"]] == ["Model!B6"]


def test_generated_code_agrees_with_the_lifted_ir():
    ir = lift_workbook(book(LOAN, {"Fees": FEES}, LOAN_NAMES), roles={"rate": "parameter"})
    X = {"B3": np.array([10.0, 20.0, 30.0]), "principal": np.array([1e5, 2e5, 3e5])}
    np.testing.assert_allclose(
        compile_reference(ir)(X, {"rate": 0.05})["B7"],
        evaluate(ir, X, {"rate": 0.05})["B7"],
        rtol=1e-14,
    )


def test_a_constant_role_carries_the_workbook_value():
    ir = lift_workbook(book(LOAN, {"Fees": FEES}, LOAN_NAMES), roles={"rate": "constant"})
    rate = next(i for i in ir["inputs"] if i["name"] == "rate")
    assert rate == {"name": "rate", "type": "float64", "role": "constant", "value": 0.06}
    with pytest.raises(ValidationFailed, match="Roles name inputs the lift does not have"):
        lift_workbook(book(LOAN, {"Fees": FEES}, LOAN_NAMES), roles={"nope": "parameter"})


# -- choosing the output ------------------------------------------------------------------------
def test_the_single_final_cell_is_found_and_ambiguity_is_refused():
    assert (
        lift_workbook(book({"A1": 2, "A2": "=A1*3"}))["lifted_from"]["workbook"]["output"]
        == "Model!A2"
    )
    two = book({"A1": 2, "A2": "=A1*3", "A3": "=A1+1"})
    with pytest.raises(ValidationFailed, match="Name the output cell.*Model!A2, Model!A3"):
        lift_workbook(two)
    assert lift_workbook(two, output="A3")["outputs"][0]["name"] == "A3"
    assert lift_workbook(two, output="Model!A2")["outputs"][0]["name"] == "A2"
    named = book({"A1": 2, "A2": "=A1*3", "A3": "=A1+1"}, names={"total": "Model!$A$3"})
    assert lift_workbook(named, output="total")["lifted_from"]["workbook"]["output"] == "Model!A3"


def test_a_cell_used_only_through_a_name_is_not_an_output_candidate():
    data = book({"A1": 2, "A2": "=A1*3", "A3": "=tripled+1"}, names={"tripled": "Model!$A$2"})
    assert lift_workbook(data)["lifted_from"]["workbook"]["output"] == "Model!A3"


@pytest.mark.parametrize(
    "output, message",
    [
        ("A1", "is a value, not a formula"),
        ("A1:A2", "not a single cell"),
        ("nosuch", "not a single cell"),
        ("Ghost!A1", "no sheet named 'Ghost'"),
    ],
)
def test_bad_outputs_are_refused(output, message):
    with pytest.raises(ValidationFailed, match=message):
        lift_workbook(book({"A1": 2, "A2": "=A1*3"}), output=output)


# -- Excel semantics, one construct at a time ------------------------------------------------------
X = {"A1": 2.0, "A2": 3.0, "A3": -1.5, "A4": 0.25}


SEMANTICS = [
    ("=A1+A2*A3", 2 + 3 * -1.5),
    ("=(A1+A2)/A4", 20.0),
    ("=-A1^2", 4.0),  # Excel: negation binds tighter than ^
    ("=2^3^2", 64.0),  # Excel: ^ is left-associative
    ("=A4*100%", 0.25),
    ("=50%", 0.5),
    ("=A1^-1", 0.5),
    ("=+A2", 3.0),
    ("=SUM(A1:A4)", 3.75),
    ("=SUM(A1:A4, 10)", 13.75),
    ("=SUM(A1)", 2.0),
    ("=PRODUCT(A1:A2)", 6.0),
    ("=MIN(A1:A4)", -1.5),
    ("=MAX(A1, A2, 7)", 7.0),
    ("=AVERAGE(A1:A4)", 3.75 / 4),
    ("=ABS(A3)", 1.5),
    ("=SQRT(A4)", 0.5),
    ("=EXP(A1)", math.exp(2)),
    ("=LN(A2)", math.log(3)),
    ("=LOG10(1000)", 3.0),
    ("=LOG(8, A1)", 3.0),
    ("=LOG(100)", 2.0),
    ("=POWER(A1, A2)", 8.0),
    ("=IF(A1>A2, 1, 2)", 2.0),
    ("=IF(A1<A2, 1)", 1.0),
    ("=IF(A1>A2, 1)", 0.0),
    ("=IF(A1<>A2, 5, 6)", 5.0),
    ("=IF(A1=2, 5, 6)", 5.0),
    ("=IF(AND(A1>0, A3<0), 1, 0)", 1.0),
    ("=IF(OR(A1<0, A3>0), 1, 0)", 0.0),
    ("=IF(NOT(A1>A2), 1, 0)", 1.0),
    ("=IF(A1>=2, 1, 0) + IF(A3<=-1.5, 1, 0)", 2.0),
    ("=NORM.S.DIST(A4, TRUE)", float(ncdf(0.25))),
    ("=_xlfn.NORM.S.DIST(A4, FALSE)", float(npdf(0.25))),
    ("=NORMSDIST(A3)", float(ncdf(-1.5))),
    ("=NORM.DIST(A2, A1, A4, TRUE)", float(ncdf(4.0))),
    ("=NORM.DIST(A2, A1, A4, FALSE)", float(npdf(4.0)) / 0.25),
    ("=TRUE + TRUE()", 2.0),
]


@pytest.mark.parametrize("formula, expected", SEMANTICS)
def test_excel_semantics_are_reproduced(formula, expected):
    ir = one(formula, X)
    assert value_of(ir) == pytest.approx(expected, rel=1e-12, abs=1e-15)
    wb = cached(book({**X, "Z1": formula}), {"Z1": expected})
    assert lift_workbook(wb, output="Z1")["lifted_from"]["workbook"]["check"]["status"] == "agreed"


def test_true_and_false_written_as_functions_are_literals():
    """LibreOffice saves the literal TRUE as TRUE(); both mean the same to Excel."""
    assert value_of(one("=NORM.S.DIST(A4, TRUE())", X)) == pytest.approx(float(ncdf(0.25)))
    table = {"A1": 1, "B1": 100, "A2": 2, "B2": 200}
    assert value_of(one("=VLOOKUP(2, A1:B2, 2, FALSE())", table)) == 200.0


def test_empty_cells_read_as_zero_and_are_skipped_in_aggregates():
    ir = one("=A1+A9", {"A1": 2})
    assert value_of(ir) == 2.0
    assert ir["lifted_from"]["workbook"]["warnings"] == [
        "Model!Z1: empty cell Model!A9 reads as 0, as in Excel"
    ]
    assert value_of(one("=AVERAGE(A1:A3)", {"A1": 2, "A3": 4})) == 3.0  # A2 skipped, not 0


def test_cross_sheet_references_are_named_by_sheet():
    ir = lift_workbook(book({"A1": "=Inputs!B2*2"}, {"Inputs": {"B2": 21}}), output="A1")
    assert [i["name"] for i in ir["inputs"]] == ["Inputs_B2"]
    assert ir["outputs"][0]["name"] == "Model_A1" and value_of(ir) == 42.0
    quoted = lift_workbook(book({"A1": "='My Inputs'!B2+1"}, {"My Inputs": {"B2": 1}}), output="A1")
    assert value_of(quoted) == 2.0 and quoted["inputs"][0]["name"] == "My_Inputs_B2"


def test_named_ranges_work_inside_aggregates_and_lookups():
    data = book(
        {
            "A1": 1,
            "A2": 2,
            "A3": 3,
            "B1": 10,
            "B2": 20,
            "B3": 30,
            "C1": 2,
            "D1": "=SUM(vals) + VLOOKUP(C1, grid, 2, FALSE)",
        },
        names={"vals": "Model!$A$1:$A$3", "grid": "Model!$A$1:$B$3"},
    )
    assert value_of(lift_workbook(data, output="D1")) == 26.0


@pytest.mark.parametrize("key, expected", [(0, math.nan), (15, 1.01), (20, 1.01), (30, 1.02)])
def test_approximate_vlookup_takes_the_largest_key_not_above(key, expected):
    ir = lift_workbook(
        book(
            {"A1": key, "Z1": "=VLOOKUP(A1, T!A1:B3, 2)"},
            {"T": {"A1": 10, "B1": 1.0, "A2": 15, "B2": 1.01, "A3": 25, "B3": 1.02}},
        ),
        output="Z1",
    )
    got = value_of(ir)
    assert (math.isnan(got) and math.isnan(expected)) or got == expected


def test_exact_lookups_and_hlookup_and_a_miss_that_is_na():
    table = {"A1": 1, "B1": 100, "A2": 2, "B2": 200, "C1": 1, "C2": 2, "D1": 5, "D2": 6}
    assert value_of(one("=VLOOKUP(2, A1:B2, 2, FALSE)", table)) == 200.0
    assert value_of(one("=HLOOKUP(5, C1:D2, 2, 0)", table)) == 6.0  # first row is the key
    miss = one("=VLOOKUP(3, A1:B2, 2, FALSE)", table)
    assert math.isnan(value_of(miss)) and r"\mathrm{N/A}" in miss["latex"]
    agreed = lift_workbook(
        cached(book({**table, "Z1": "=VLOOKUP(3, A1:B2, 2, FALSE)"}), {"Z1": "#N/A"}), output="Z1"
    )
    assert agreed["lifted_from"]["workbook"]["check"]["status"] == "agreed"


# -- refusals, named by cell ---------------------------------------------------------------------
@pytest.mark.parametrize(
    "formula, cells, message",
    [
        ('="a"&"b"', {}, "text literal"),
        ("=A1&A2", {"A1": 1, "A2": 2}, "concatenation"),
        ("=ROUND(A1, 2)", {"A1": 1}, "the function ROUND"),
        ("=NPV(0.1, A1:A2)", {"A1": 1, "A2": 2}, "the function NPV"),
        ("=TODAY()", {}, "the function TODAY"),
        ("=INDIRECT(A1)", {"A1": 1}, "the function INDIRECT"),
        ("=INDEX(A1:A2, 1)", {"A1": 1, "A2": 2}, "the function INDEX"),
        ("=A1+1", {"A1": "text"}, "holding text"),
        ("=A1:A2*2", {"A1": 1, "A2": 2}, "outside an aggregate or lookup"),
        ("=#N/A", {}, "error literal"),
        ("=nosuch*2", {}, "no defined name 'nosuch'"),
        ("=Ghost!A1", {}, "no sheet named 'Ghost'"),
        (
            "=VLOOKUP(1, A1:B2, A3)",
            {"A1": 1, "B1": 2, "A2": 3, "B2": 4, "A3": 2},
            "column index that is not a literal",
        ),
        ("=VLOOKUP(1, A1:B2, 5)", {"A1": 1, "B1": 2, "A2": 3, "B2": 4}, "outside the table"),
        ("=VLOOKUP(1, A1:B2, 2)", {"A1": 3, "B1": 2, "A2": 1, "B2": 4}, "strictly ascending"),
        (
            "=VLOOKUP(1, A1:B2, 2, FALSE)",
            {"A1": 1, "B1": "=2*2", "A2": 3, "B2": 4},
            "table containing formulas",
        ),
        (
            "=VLOOKUP(1, A1:B2, 2, A3)",
            {"A1": 1, "B1": 2, "A2": 3, "B2": 4, "A3": 0},
            "not written literally",
        ),
        ("=NORM.S.DIST(1)", {}, "takes 2 argument"),
        ("=SQRT(1, 2)", {}, "takes 1 argument"),
        ("=IF(1)", {}, "takes 2 to 3"),
        ("=SUM(A1:A2)", {}, "an aggregate over no numbers"),
        ("=(1+2", {}, "expected"),
        ("=1+2)", {}, "unexpected"),
        ("=A:A", {}, "unexpected ':'"),  # whole-column references are not parsed
    ],
)
def test_what_is_out_of_scope_is_refused_by_cell(formula, cells, message):
    with pytest.raises(ValidationFailed, match=message) as info:
        one(formula, cells)
    assert "Model!Z1" in info.value.message or "Model!" in str(info.value.context)


def test_circular_references_are_refused_with_the_cycle():
    with pytest.raises(
        ValidationFailed,
        match=r"circular reference: Model!A1 -> Model!A2 -> "
        r"Model!A1",
    ):
        lift_workbook(book({"A1": "=A2+1", "A2": "=A1*2", "A3": "=A1"}), output="A3")


def test_unreadable_or_oversized_files_are_refused(monkeypatch):
    with pytest.raises(ValidationFailed, match="Not a readable .xlsx"):
        lift_workbook(b"PK\x03\x04 not really a workbook")
    import maya.formula.xlsx as xl

    monkeypatch.setattr(xl, "MAX_BYTES", 10)
    with pytest.raises(ValidationFailed, match="over 0 MB"):
        lift_workbook(book({"A1": 1, "A2": "=A1"}))


def test_the_tokenizer_tells_functions_from_cell_like_names():
    assert tokenize("LOG10(A1)", "x")[0] == ("name", "LOG10")
    assert tokenize("'My Sheet'!$B$2:C3", "x") == [("ref", "'My Sheet'!$B$2:C3")]
    assert [k for k, _ in tokenize('1.5e3 <> "a""b"', "x")] == ["num", "op", "str"]


# -- the platform: service, API, SDK, CLI, check ---------------------------------------------------
@pytest.fixture(scope="module")
def sheets(world):
    from maya.api.app import create_api
    from maya.sdk import Client

    w = world
    w.p.access.create_namespace(w.admin, name="xl", preset="standard")
    app = create_api(w.p)
    mona = Client(app=app, token=Client(app=app).auth.login("mona", "Test-password-1")["token"])
    return w, mona


def test_import_into_a_draft_keeps_the_workbook_and_audits(sheets):
    w, mona = sheets
    data = cached(book(LOAN, {"Fees": FEES}, LOAN_NAMES), loan_values())
    mona.models.create("xl", "loan", description="from Excel")
    out = mona.models.import_workbook(
        "xl/loan", data, roles={"rate": "parameter"}, filename="loan.xlsx"
    )
    assert out["workbook"]["check"]["status"] == "agreed" and out["workbook"]["blob"]
    got = mona.models.get("xl/loan")["versions"][0]
    assert got["formula_ir"]["outputs"] == [{"name": "B7", "type": "float64"}]
    assert [c["name"] for c in got["input_contract"]] == ["B3", "principal"]
    assert mona.models.workbook("xl/loan", 1)["data"] == data
    with w.p.uow() as uow:
        entry = uow.repo("audit_events").list(action="model.workbook_imported")[-1]
    assert entry["detail"]["check"] == "agreed" and entry["detail"]["output"] == "Model!B7"


def test_preview_stores_nothing(sheets):
    w, mona = sheets
    ir = mona.models.lift_workbook(book({"A1": 2, "A2": "=A1*3"}), filename="tiny.xlsx")
    assert (
        ir["lifted_from"]["workbook"]["filename"] == "tiny.xlsx"
        and "blob" not in ir["lifted_from"]["workbook"]
    )
    from maya.core.errors import ValidationFailed as VF

    with pytest.raises(VF, match="roles is not valid JSON|Roles name"):
        mona.models.lift_workbook(book({"A1": 2, "A2": "=A1*3"}), roles={"ghost": "parameter"})


def test_a_disagreeing_workbook_cannot_be_submitted(sheets):
    w, mona = sheets
    wrong = {**loan_values(), "B7": 1.0}
    mona.models.create("xl", "badloan")
    mona.models.import_workbook("xl/badloan", cached(book(LOAN, {"Fees": FEES}), wrong))
    ok, why = w.p.models.check_formula(
        None, {"row": w.p.models.get(w.mona, "xl/badloan")["versions"][0]}
    )
    assert not ok and "disagrees with the workbook" in why


def test_a_model_not_lifted_from_a_workbook_has_none_to_download(sheets):
    from maya.core.errors import NotFound

    w, mona = sheets
    mona.models.create("xl", "plain", formula="y = a*x", roles={"a": "parameter"})
    with pytest.raises(NotFound, match="not lifted from a workbook"):
        mona.models.workbook("xl/plain", 1)


def test_the_cli_previews_and_imports(sheets, tmp_path, monkeypatch, capsys):
    import maya.cli.__main__ as cli
    from maya.cli import common as cli_common

    w, mona = sheets
    # every command group opens its client through this one function
    monkeypatch.setattr(cli_common, "client", lambda args: mona)
    path = tmp_path / "loan.xlsx"
    path.write_bytes(cached(book(LOAN, {"Fees": FEES}, LOAN_NAMES), loan_values()))
    assert cli.main(["model", "import-workbook", "xl/unused", str(path), "--preview"]) == 0
    assert "output Model!B7" in capsys.readouterr().out
    mona.models.create("xl", "cliloan")
    assert (
        cli.main(
            [
                "--json",
                "model",
                "import-workbook",
                "xl/cliloan",
                str(path),
                "--role",
                "rate=parameter",
            ]
        )
        == 0
    )
    assert '"status": "agreed"' in capsys.readouterr().out
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(cached(book(LOAN, {"Fees": FEES}), {**loan_values(), "B7": 1.0}))
    mona.models.create("xl", "clibad")
    assert cli.main(["model", "import-workbook", "xl/clibad", str(bad)]) == 1


def test_the_web_form_creates_a_model_from_a_workbook_and_shows_the_check(sheets):
    """New model → 'Lift from an Excel workbook'; the Definition tab shows the check; the
    original downloads; a re-import into the draft works; an empty upload is refused."""
    from starlette.testclient import TestClient

    from maya.server import build_app

    w, _ = sheets
    with w.p.uow() as uow:  # not under test here: skip the first-login change
        uow.repo("users").update(
            uow.repo("users").find_one(username="mona")["id"], {"must_change_password": False}
        )
    web = TestClient(build_app(w.p))
    csrf = re.compile(r'name="csrf_token" value="([^"]+)"')
    token = csrf.search(web.get("/login").text).group(1)
    web.post(
        "/login", data={"username": "mona", "password": "Test-password-1", "csrf_token": token}
    )
    page = web.get("/models/new").text
    assert "<h1>New model</h1>" in page, re.findall(r"<h1[^>]*>([^<]*)", page)
    assert 'value="workbook"' in page and 'enctype="multipart/form-data"' in page
    data = cached(book(LOAN, {"Fees": FEES}, LOAN_NAMES), loan_values())
    form = {
        "csrf_token": csrf.search(page).group(1),
        "namespace": "xl",
        "name": "webloan",
        "kind": "formula",
        "authoring": "workbook",
        "workbook_roles": "rate: parameter",
    }
    r = web.post("/models/new", data=form, files={"workbook": ("loan.xlsx", data)})
    assert r.status_code == 200 and "Model created from loan.xlsx" in r.text
    assert "Lifted from loan.xlsx" in r.text and "Agreed." in r.text
    got = web.get("/models/xl/webloan/versions/1/workbook.xlsx")
    assert got.status_code == 200 and got.content == data
    wrong = cached(book(LOAN, {"Fees": FEES}), {**loan_values(), "B7": 1.0})
    tok = csrf.search(web.get("/models/xl/webloan?tab=definition").text).group(1)
    r = web.post(
        "/models/xl/webloan/workbook",
        data={"csrf_token": tok},
        files={"workbook": ("bad.xlsx", wrong)},
    )
    assert "Disagreed." in r.text and "Model!B7" in r.text
    tok = csrf.search(web.get("/models/new").text).group(1)
    r = web.post(
        "/models/new",
        data={**form, "csrf_token": tok, "name": "empty"},
        files={"workbook": ("", b"")},
    )
    assert "Choose an .xlsx workbook" in r.text
