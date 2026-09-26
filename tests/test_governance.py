"""
Model governance: the findings register, materiality tiers and periodic review.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
"""

from __future__ import annotations

import datetime as dt

import pytest

from maya.core.clock import utcnow
from maya.core.errors import NotApproved, PermissionDenied, ValidationFailed, WarrantSuspended


@pytest.fixture(scope="module")
def gov(world):
    w = world
    w.p.models.create(
        w.mona,
        namespace="eq",
        name="gov_lin",
        formula="yhat = a*x + b",
        roles={"a": "parameter", "b": "parameter"},
    )
    return w


def _live_warrant(w, name: str, approved_days_ago: int) -> str:
    """A model version approved long ago and one sealed, live execution warrant on it.

    Built directly: the warrant chain has its own tests, and what is under test here is
    what governance does to a warrant that is live."""
    with w.p.uow("test") as uow:
        model = uow.repo("models").find_one(name="gov_lin")
        v = uow.repo("model_versions").find_one(model_id=model["id"], version_no=1)
        uow.repo("model_versions").update(
            v["id"], {"approved_at": utcnow() - dt.timedelta(days=approved_days_ago)}
        )
        ew = uow.repo("execution_warrants").add(
            {
                "namespace_id": model["namespace_id"],
                "name": name,
                "state": "approved",
                "owner_id": model["owner_id"],
                "model_version_id": v["id"],
                "spec": {"environments": ["dev"], "contact": "risk@example.com"},
                "sealed_at": utcnow(),
                "valid_from": utcnow(),
                "valid_to": utcnow() + dt.timedelta(days=90),
            }
        )
        return ew["id"]


def test_a_finding_is_raised_remediated_and_closed_by_someone_independent(gov):
    w = gov
    f = w.p.governance.raise_finding(
        w.devi, "eq/gov_lin", "Intercept unstable across samples", "high", detail="see memo"
    )
    assert f["state"] == "open" and f["owner"] == "mona"
    assert f["due_date"] == utcnow().date() + dt.timedelta(days=90)
    w.p.governance.move_finding(w.mona, f["id"], "start")
    w.p.governance.move_finding(w.mona, f["id"], "remediated")
    with pytest.raises(PermissionDenied, match="does not close"):
        w.p.governance.move_finding(w.mona, f["id"], "close", "fixed it")
    with pytest.raises(ValidationFailed, match="note"):
        w.p.governance.move_finding(w.devi, f["id"], "close")
    done = w.p.governance.move_finding(w.devi, f["id"], "close", "re-ran; stable now")
    assert done["state"] == "closed" and done["closed_by"] == "devi"
    assert [h["action"] for h in done["history"]] == ["raised", "start", "remediated", "close"]
    with pytest.raises(NotApproved):
        w.p.governance.move_finding(w.devi, f["id"], "start")
    assert f["id"] not in {r["id"] for r in w.p.governance.findings(w.devi, state="active")}


def test_the_owner_does_not_accept_the_risk_in_their_own_model(gov):
    w = gov
    f = w.p.governance.raise_finding(w.devi, "eq/gov_lin", "No challenger", "low")
    with pytest.raises(PermissionDenied):
        w.p.governance.move_finding(w.mona, f["id"], "accept", "fine")
    out = w.p.governance.move_finding(w.admin, f["id"], "accept", "compensating control")
    assert out["state"] == "accepted" and out["resolution"] == "compensating control"


def test_the_tier_is_derived_and_an_override_that_lowers_it_is_flagged(gov):
    w = gov
    base = w.p.governance.profile(w.mona, "eq/gov_lin")
    assert base["derived_tier"] == 3 and not base["declared"]
    prof = w.p.governance.set_profile(w.mona, "eq/gov_lin", use="regulatory", exposure=5e9)
    assert prof["derived_tier"] == 1 and prof["tier"] == 1 and prof["review_days"] == 365
    with pytest.raises(ValidationFailed, match="reason"):
        w.p.governance.set_profile(w.mona, "eq/gov_lin", use="regulatory", tier_override=3)
    low = w.p.governance.set_profile(
        w.mona,
        "eq/gov_lin",
        use="regulatory",
        exposure=5e9,
        tier_override=3,
        override_reason="shadow use only",
    )
    assert low["tier"] == 3 and low["override_lowers"]
    with pytest.raises(PermissionDenied):
        w.p.governance.set_profile(w.devi, "eq/gov_lin", use="internal")
    w.p.governance.set_profile(w.mona, "eq/gov_lin", use="regulatory", exposure=5e9)


def test_an_overdue_review_suspends_live_warrants_and_a_review_lifts_only_those(gov):
    w = gov
    ew_id = _live_warrant(w, "gov_live", approved_days_ago=400)  # tier 1: yearly
    other = _live_warrant(w, "gov_breached", approved_days_ago=400)
    with w.p.uow("test") as uow:  # suspended for something else entirely
        uow.repo("execution_warrants").update(
            other, {"suspended_at": utcnow(), "suspend_reason": "covenant breach"}
        )
    assert w.p.governance.profile(w.mona, "eq/gov_lin")["review_overdue"]
    with pytest.raises(PermissionDenied):
        w.p.governance.sweep(w.devi)
    out = w.p.governance.sweep(w.admin)
    assert out["suspended"] == [ew_id]
    with w.p.uow() as uow:
        ew = uow.repo("execution_warrants").require(ew_id)
    with pytest.raises(WarrantSuspended, match="Periodic review overdue"):
        w.p.execution.check(ew)
    with pytest.raises(PermissionDenied, match="owner"):
        w.p.governance.record_review(w.mona, "eq/gov_lin", "satisfactory", "all good")
    rec = w.p.governance.record_review(w.devi, "eq/gov_lin", "satisfactory", "backtest reviewed")
    assert rec["reinstated"] == [ew_id] and not rec["review_overdue"]
    assert rec["next_review_due"] == utcnow().date() + dt.timedelta(days=365)
    with w.p.uow() as uow:
        assert uow.repo("execution_warrants").require(ew_id)["suspended_at"] is None
        assert uow.repo("execution_warrants").require(other)["suspend_reason"] == "covenant breach"
    assert w.p.governance.sweep(w.admin)["suspended"] == []
    row = next(r for r in w.p.governance.overview(w.mona)["models"] if r["name"] == "gov_lin")
    assert row["tier"] == 1 and not row["review_overdue"]


def test_the_register_is_reachable_through_the_sdk(gov):
    from maya.sdk.client import Client

    assert hasattr(Client, "__init__")
    from maya.sdk.resources import ENDPOINTS

    assert ENDPOINTS[("POST", "/governance/findings")] == "Governance.raise_finding"


# ------------------------------------------------------------------ monitoring


def _reports(w, ew_id: str, runs: list[dict]) -> None:
    with w.p.uow("test") as uow:
        for i, r in enumerate(runs):
            row = uow.repo("execution_reports").add(
                {"execution_warrant_id": ew_id, "environment": "dev", **r}
            )
            when = utcnow() - dt.timedelta(days=len(runs) - i)
            uow.repo("execution_reports").update(row["id"], {"created_at": when})


def test_monitoring_reads_reports_as_series_and_grades_the_warrant(gov):
    w = gov
    ew_id = _live_warrant(w, "gov_monitored", approved_days_ago=10)
    baseline = [10, 10, 10, 10]
    with w.p.uow("test") as uow:
        uow.repo("execution_warrants").update(
            ew_id,
            {
                "spec": {
                    "environments": ["dev"],
                    "covenants": [
                        {
                            "kind": "input_psi",
                            "attr": "x",
                            "max": 0.25,
                            "baseline": {"edges": [0, 1, 2, 3, 4], "counts": baseline},
                        },
                        {"kind": "input_null_rate", "attr": "x", "max": 0.5},
                    ],
                }
            },
        )
    steady = {"null_rate": 0.01, "mean": 1.5, "min": 0.0, "max": 3.9, "histogram": baseline}
    drifted = {"null_rate": 0.05, "mean": 2.4, "min": 0.0, "max": 3.9, "histogram": [4, 8, 12, 16]}
    runs = [
        {"rows": 100, "input_stats": {"x": steady}, "output_stats": {"yhat": {"mean": 3.0}}}
        for _ in range(5)
    ]
    runs.append(
        {"rows": 120, "input_stats": {"x": drifted}, "output_stats": {"yhat": {"mean": 4.0}}}
    )
    _reports(w, ew_id, runs)

    m = w.p.monitoring.warrant(w.mona, ew_id, days=30)
    xs = m["series"]["inputs"]["x"]
    assert len(xs) == 6 and xs[0]["psi"] == 0.0 and 0.10 < xs[-1]["psi"] < 0.25
    assert m["rows"] == 620 and m["runs"] == 6 and len(m["daily"]) == 6
    assert m["bounds"]["x"]["null_max"] == 0.5
    levels = {(s["what"], s["level"]) for s in m["signals"]}
    assert ("x", "watch") in levels  # PSI in the watch band and a doubled null rate
    assert not any(s["level"] == "breach" for s in m["signals"])

    row = next(r for r in w.p.monitoring.overview(w.mona)["warrants"] if r["id"] == ew_id)
    assert row["health"] == "watch" and row["worst_psi"] == xs[-1]["psi"]
    with pytest.raises(PermissionDenied):
        w.p.monitoring.warrant(w.principal("dana"), ew_id)


def test_chart_geometry_keeps_bounds_on_the_scale():
    from maya.web.charts import line_chart, nice_ticks

    assert nice_ticks(0.0, 0.23) == [0.0, 0.1, 0.2, 0.3]
    pts = [{"at": f"2026-03-0{i}T00:00:00", "v": v} for i, v in enumerate([0.1, 0.2, 0.15], 1)]
    c = line_chart(pts, "v", bounds=[(0.5, "covenant")], marks=["2026-03-02T00:00:00"], floor=0)
    assert not c["empty"] and c["last"] == 0.15
    # the axis reaches the covenant rather than cropping it, and draws it inside the plot
    assert float(c["yticks"][-1][1]) >= 0.5
    assert c["top"] <= c["lines"][0][0] < c["bottom"] and len(c["marks"]) == 1
    assert line_chart([], "v")["empty"]


# ------------------------------------------------------------------- inventory


def test_the_inventory_exports_in_each_layout_and_counts_what_it_leaves_out(gov):
    import csv
    import io
    import json

    from openpyxl import load_workbook

    w = gov
    out = w.p.inventory.export(w.mona, "json", "ss1-23")
    doc = json.loads(out["data"])
    row = next(r for r in doc["rows"] if r["name"] == "gov_lin")
    assert row["tier"] == 1 and row["use"] == "regulatory" and row["owner"] == "mona"
    assert row["last_review_outcome"] == "satisfactory" and row["review_overdue"] == "no"
    assert row["accepted_risks"] >= 1 and "use: regulatory" in row["tier_basis"]
    assert doc["columns"]["tier"] == "Model tier" and "restrictions" in doc["columns"]
    assert "not a regulatory submission template" in doc["notice"]

    sr = w.p.inventory.export(w.mona, "csv", "sr11-7")
    lines = list(csv.reader(io.StringIO(sr["data"].decode())))
    assert lines[0][0].startswith("# Model inventory, SR 11-7 layout")
    header = lines[2]
    assert "Risk rating" in header and "Restrictions on use" not in header  # SS1/23 only
    assert any(r and r[0] == "maya://model/eq/gov_lin" for r in lines[3:])

    xl = w.p.inventory.export(w.mona, "xlsx", "sr11-7")
    wb = load_workbook(io.BytesIO(xl["data"]))
    assert wb.sheetnames == ["Inventory", "About"] and wb["Inventory"]["A1"].value == "Model ID"
    assert xl["filename"].endswith(".xlsx")

    # a model in a namespace devi is not staffed in is left out of her inventory, and counted
    w.p.access.create_namespace(w.admin, name="tiny", preset="small_team")
    w.p.models.create(w.admin, namespace="tiny", name="gov_hidden", formula="y = 2*x")
    seen = json.loads(w.p.inventory.export(w.devi, "json", "maya")["data"])
    assert "gov_hidden" not in {r["name"] for r in seen["rows"]} and seen["not_shown"] >= 1
    everything = json.loads(w.p.inventory.export(w.admin, "json", "maya")["data"])
    assert "gov_hidden" in {r["name"] for r in everything["rows"]}
    with pytest.raises(ValidationFailed):
        w.p.inventory.export(w.mona, "pdf")
