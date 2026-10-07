"""
The storytelling layouts: the slide kinds a briefing deck is told with.

``tldr`` (question-and-answer rows), ``numbered`` (numbered rows: principles, steps,
a worked example, with an optional question above and a key insight below),
``quotes``, ``arch`` (layered bands), ``lanes`` (two teams meeting at one contract),
``columns`` (side-by-side options), ``compare`` (a capability matrix whose cells are
graded by colour), ``shot`` (a real screenshot with its commentary), ``cover`` and
``thanks``. Each is drawn from a plain dictionary, like the kinds in ``layouts``, and
every block of text is fitted with the estimator the geometry audit uses.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
Proprietary and confidential. See LICENSE and NOTICE at the repository root.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as In

import theme as T
from metrics import SAFETY, SANS, SH, SW, text_h

from maya.core.version import VERSION

REPO = Path(__file__).resolve().parents[2]
SCREENS = REPO / "maya" / "web" / "static" / "help" / "screens"
GAP = 0.16


def _round(sl: Any, x: float, y: float, w: float, h: float, fill: Any, line: Any = None) -> Any:
    s = T.rect(sl, x, y, w, h, fill=fill, line=line, lw=0.9, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    s.adjustments[0] = min(0.5, 0.10 / max(0.2, min(w, h)))
    return s


def _size_for(cells: list[tuple[str, float, float, bool]], start: float, floor: float) -> float:
    """The largest size at which every ``(text, width, height, bold)`` cell fits."""
    size = start
    while size > floor:
        if all(text_h(t, w * SAFETY, size, b, SANS, 1.15) <= h for t, w, h, b in cells):
            return size
        size -= 0.5
    return floor


def _cell(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: float,
    color: Any,
    bold: bool = False,
    italic: bool = False,
    align: Any = PP_ALIGN.LEFT,
) -> None:
    """One block of text, vertically centred in its box, shrunk further if it must."""

    def write(tf: Any, s: float) -> None:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = align
        T.para(
            tf,
            text,
            size=s,
            color=color,
            bold=bold,
            italic=italic,
            first=True,
            space_after=0,
            line=1.15,
        )

    T.fitted(sl, x, y, w, h, write, size, 7.5)
    # fitted() sized the box to the text's height at the top; stretch it back to the row
    # so the middle anchor centres it.
    sl.shapes[-1].height = In(h)


def _band(sl: Any, text: str, bottom: float, label: str = "Key insight") -> float:
    """A tinted band pinned above the footer; returns its top."""
    width = T.CW - 0.5
    size = 13.0
    full = f"{label}: {text}" if label else text
    while size > 9.5 and text_h(full, width * SAFETY, size, False, SANS, 1.15) > 0.62:
        size -= 0.5
    h = text_h(full, width * SAFETY, size, False, SANS, 1.15) + 0.22
    top = bottom - h
    _round(sl, T.ML, top, T.CW, h, T.PINK_L)

    def write(tf: Any, s: float) -> None:
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        parts = [(f"{label}: ", T.CRIMSON_D, True)] if label else []
        T.runs(
            tf, parts + [(text, T.CRIMSON_D, False)], size=s, first=True, space_after=0, line=1.15
        )
        for r in tf.paragraphs[0].runs:
            r.font.italic = True

    T.fitted(sl, T.ML + 0.25, top + 0.06, width, h - 0.12, write, size, 8.5)
    return top - GAP


def _question(sl: Any, y: float, text: str | None) -> float:
    if not text:
        return y
    used = T.fitted(
        sl,
        T.ML,
        y,
        T.CW,
        0.7,
        lambda tf, s: T.para(
            tf, text, size=s, color=T.CRIMSON, italic=True, first=True, space_after=0, line=1.15
        ),
        15,
        11,
    )
    return y + used + GAP


# -- the kinds -------------------------------------------------------------------------------


def cover(s: dict[str, Any]) -> None:
    """The opening slide: crimson, the mark, the name, the slogan, the author."""
    T._state["n"] = 1
    sl = T.blank()
    T.rect(sl, 0, 0, SW, SH, fill=T.CRIMSON)
    T.rect(sl, 0, 0, 0.22, SH, fill=T.CRIMSON_D)
    T.mark(sl, T.ML, 0.75, 0.62)
    tf = T.txt(sl, T.ML + 0.82, 0.86, 4.0, 0.42)
    T.para(tf, "M A Y A", size=22, color=T.WHITE, bold=True, first=True, space_after=0)

    def head(tf: Any, size: float) -> None:
        for i, line in enumerate(s["title"]):
            T.para(
                tf,
                line,
                size=size,
                color=T.WHITE,
                bold=True,
                first=i == 0,
                space_after=0,
                line=1.05,
            )

    T.fitted(sl, T.ML, 2.0, T.CW * 0.62, 1.6, head, 40, 26)
    T.rect(sl, T.ML, 3.72, 1.6, 0.04, fill=T.GOLD)
    tf = T.txt(sl, T.ML, 3.9, T.CW * 0.62, 0.6)
    T.para(tf, s["sub"], size=24, color=T.PINK_L, italic=True, first=True, space_after=0)
    tf = T.txt(sl, T.ML, 5.15, T.CW * 0.5, 1.2)
    T.para(tf, "Ashutosh Sinha", size=18, color=T.WHITE, bold=True, first=True, space_after=3)
    T.para(tf, s.get("date", "October 2026"), size=12, color=T.PINK_L, space_after=1)
    T.para(tf, f"MAYA {VERSION}", size=11, color=T.PINK, space_after=0)
    # The chain the deck is about, down the right-hand side.
    x, w, h, gap = SW - T.ML - 3.1, 3.1, 0.62, 0.2
    top = (SH - (len(s["chain"]) * (h + gap) - gap)) / 2
    for i, (name, what) in enumerate(s["chain"]):
        y = top + i * (h + gap)
        _round(sl, x, y, w, h, T.CRIMSON_D, line=T.PINK)
        tf = T.txt(sl, x + 0.18, y + 0.08, w - 0.3, h - 0.14, anchor=MSO_ANCHOR.MIDDLE)
        T.runs(
            tf,
            [(name, T.WHITE, True), (f"   {what}", T.PINK_L, False)],
            size=11.5,
            first=True,
            space_after=0,
        )
        if i:
            T.connect(sl, x + w / 2, y - gap + 0.03, x + w / 2, y - 0.03, T.PINK, 1.5)


def thanks(s: dict[str, Any]) -> None:
    T._state["n"] += 1
    sl = T.blank()
    T.rect(sl, 0, 0, SW, SH, fill=T.CRIMSON)
    T.rect(sl, 0, 0, 0.22, SH, fill=T.CRIMSON_D)
    T.mark(sl, (SW - 0.9) / 2, 1.2, 0.9)
    tf = T.txt(sl, T.ML, 2.45, T.CW, 0.9, align=PP_ALIGN.CENTER)
    T.para(
        tf,
        s.get("title", "Thank you"),
        size=44,
        color=T.WHITE,
        bold=True,
        first=True,
        space_after=0,
    )
    tf = T.txt(sl, T.ML, 3.5, T.CW, 0.5, align=PP_ALIGN.CENTER)
    T.para(tf, s["sub"], size=20, color=T.PINK_L, italic=True, first=True, space_after=0)
    tf = T.txt(sl, T.ML, 4.6, T.CW, 1.2, align=PP_ALIGN.CENTER)
    for i, line in enumerate(s["lines"]):
        p = T.para(
            tf,
            line,
            size=15 if i == 0 else 12,
            color=T.WHITE if i == 0 else T.PINK_L,
            bold=i == 0,
            first=i == 0,
            space_after=4,
        )
        p.alignment = PP_ALIGN.CENTER


def tldr(s: dict[str, Any]) -> None:
    """Question-and-answer rows: the executive summary."""
    sl, y = T.content(s["title"])
    bottom = T.BODY_BOTTOM
    rows = s["rows"]
    rh = (bottom - y - GAP * (len(rows) - 1)) / len(rows)
    qw, aw = T.CW * 0.27, T.CW * 0.73 - 0.75
    qs = _size_for([(q, qw, rh - 0.14, True) for q, _ in rows], 16, 10)
    an = _size_for([(a, aw, rh - 0.14, False) for _, a in rows], 13.5, 8.5)
    for i, (q, a) in enumerate(rows):
        top = y + i * (rh + GAP)
        _round(sl, T.ML, top, T.CW, rh, T.PINK_L if i % 2 == 0 else T.PARCH)
        _cell(sl, T.ML + 0.25, top + 0.06, qw, rh - 0.12, q, qs, T.CRIMSON_D, bold=True)
        _cell(sl, T.ML + 0.25 + qw + 0.25, top + 0.06, aw, rh - 0.12, a, an, T.INK)


def numbered(s: dict[str, Any]) -> None:
    """Numbered rows -- principles, steps, a worked example -- each a disc, a head and a line.

    ``question`` puts an italic prompt above; ``insight`` pins a band below."""
    sl, y = T.content(s["title"])
    y = _question(sl, y, s.get("question"))
    bottom = _band(sl, s["insight"], T.BODY_BOTTOM) if s.get("insight") else T.BODY_BOTTOM
    items = s["items"]
    n = len(items)
    rh = (bottom - y) / n
    d = min(0.46, rh - 0.08)
    hx = T.ML + d + 0.25
    hw = T.CW * s.get("head_w", 0.27)
    bx = hx + hw + 0.25
    bw = T.ML + T.CW - bx
    hs = _size_for([(h, hw, rh - 0.06, True) for h, _ in items], s.get("size", 15), 9.5)
    bs = _size_for([(b, bw, rh - 0.06, False) for _, b in items], s.get("size", 15) - 1.5, 8.5)
    for i, (head, body) in enumerate(items):
        top = y + i * rh
        cy = top + (rh - d) / 2
        T.rect(sl, T.ML, cy, d, d, fill=T.CRIMSON, shape=MSO_SHAPE.OVAL)
        tf = T.txt(sl, T.ML, cy, d, d, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T.para(
            tf,
            str(i + 1),
            size=min(14, d * 30),
            color=T.WHITE,
            bold=True,
            first=True,
            space_after=0,
            line=1.0,
        )
        _cell(sl, hx, top + 0.03, hw, rh - 0.06, head, hs, T.CRIMSON_D, bold=True)
        _cell(sl, bx, top + 0.03, bw, rh - 0.06, body, bs, T.SLATE)


def quotes(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"])
    bottom = _band(sl, s["insight"], T.BODY_BOTTOM, "") if s.get("insight") else T.BODY_BOTTOM
    items = s["items"]
    rh = (bottom - y - GAP * (len(items) - 1)) / len(items)
    w = T.CW - 0.55
    qs = _size_for([(q, w, rh - 0.5, False) for q, _ in items], 15, 10)
    for i, (quote, who) in enumerate(items):
        top = y + i * (rh + GAP)
        _round(sl, T.ML, top, T.CW, rh, T.PINK_L)
        T.rect(sl, T.ML, top + 0.12, 0.06, rh - 0.24, fill=T.CRIMSON)

        def write(tf: Any, size: float, q: str = quote, wh: str = who) -> None:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            T.para(
                tf,
                f"“{q}”",
                size=size,
                color=T.INK,
                italic=True,
                first=True,
                space_after=4,
                line=1.15,
            )
            T.para(
                tf, f"— {wh}", size=max(8.5, size - 3), color=T.CRIMSON, bold=True, space_after=0
            )

        T.fitted(sl, T.ML + 0.3, top + 0.08, w, rh - 0.16, write, qs, 8.5)
        sl.shapes[-1].height = In(rh - 0.16)


def arch(s: dict[str, Any]) -> None:
    """Layered bands, top to bottom, then optional items beneath."""
    sl, y = T.content(s["title"])
    bottom = T.BODY_BOTTOM
    bands = s["bands"]
    bh, gap = s.get("band_h", 0.5), 0.09
    fills = [T.CRIMSON_D, T.CRIMSON, T.ROSE, T.NAVY, T.SLATE, T.MUTED]
    lw = 2.1
    for i, (label, body) in enumerate(bands):
        top = y + i * (bh + gap)
        _round(sl, T.ML, top, T.CW, bh, fills[i % len(fills)])
        _cell(sl, T.ML + 0.2, top + 0.04, lw, bh - 0.08, label, 13, T.WHITE, bold=True)
        _cell(
            sl, T.ML + 0.2 + lw + 0.15, top + 0.04, T.CW - lw - 0.55, bh - 0.08, body, 12.5, T.WHITE
        )
    top = y + len(bands) * (bh + gap) + GAP
    if s.get("items"):
        _bullets(sl, top, bottom, s["items"], s.get("size", 13.5))


def _bullets(sl: Any, top: float, bottom: float, items: list[str], start: float) -> None:
    def write(tf: Any, size: float) -> None:
        for i, it in enumerate(items):
            T.runs(
                tf,
                [("▪  ", T.CRIMSON, True), (it, T.SLATE, False)],
                size=size,
                first=i == 0,
                space_after=0,
                space_before=0 if i == 0 else 5,
            )

    T.fitted(sl, T.ML, top, T.CW, bottom - top, write, start, 8.5)


def lanes(s: dict[str, Any]) -> None:
    """Two teams, each a row of steps, meeting at one contract in the middle."""
    sl, y = T.content(s["title"])
    lane_h, band_h, gap = 1.12, 0.52, 0.26
    label_w, arrow = 1.75, 0.26
    rows = [(s["top"], y), (s["bottom"], y + lane_h + gap + band_h + gap)]
    for lane, ly in rows:
        _round(sl, T.ML, ly, label_w, lane_h, T.PINK_L)
        _cell(
            sl,
            T.ML + 0.12,
            ly + 0.06,
            label_w - 0.24,
            lane_h - 0.12,
            lane["name"],
            14,
            T.CRIMSON_D,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        steps = lane["steps"]
        x0 = T.ML + label_w + 0.3
        bw = (T.ML + T.CW - x0 - arrow * (len(steps) - 1)) / len(steps)
        for i, (head, body) in enumerate(steps):
            bx = x0 + i * (bw + arrow)
            T.card(sl, bx, ly, bw, lane_h, "", head, body)
            if i:
                T.connect(
                    sl,
                    bx - arrow + 0.03,
                    ly + lane_h / 2,
                    bx - 0.03,
                    ly + lane_h / 2,
                    T.CRIMSON,
                    2.0,
                )
    by = y + lane_h + gap
    _round(sl, T.ML + label_w + 0.3, by, T.CW - label_w - 0.3, band_h, T.CRIMSON)
    _cell(
        sl,
        T.ML + label_w + 0.5,
        by + 0.04,
        T.CW - label_w - 0.7,
        band_h - 0.08,
        s["middle"],
        13.5,
        T.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    top = rows[1][1] + lane_h + GAP + 0.04
    if s.get("items"):
        _bullets(sl, top, T.BODY_BOTTOM, s["items"], s.get("size", 13))


def columns(s: dict[str, Any]) -> None:
    """Side-by-side options, each a coloured head over a tinted body."""
    sl, y = T.content(s["title"])
    cols = s["cols"]
    ch = s.get("col_h", 2.4)
    cw = (T.CW - GAP * 1.5 * (len(cols) - 1)) / len(cols)
    for i, col in enumerate(cols):
        x = T.ML + i * (cw + GAP * 1.5)
        _round(sl, x, y, cw, 0.48, T.CRIMSON)
        _cell(
            sl,
            x + 0.12,
            y + 0.04,
            cw - 0.24,
            0.4,
            col["head"],
            14,
            T.WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        body_top = y + 0.56
        _round(sl, x, body_top, cw, ch - 0.56, T.PARCH)
        lines = col["items"]
        foot = col.get("foot")

        def write(tf: Any, size: float, lines: list[str] = lines, foot: str | None = foot) -> None:
            for j, line in enumerate(lines):
                T.runs(
                    tf,
                    [("▪  ", T.CRIMSON, True), (line, T.INK, False)],
                    size=size,
                    first=j == 0,
                    space_after=0,
                    space_before=0 if j == 0 else 4,
                )
            if foot:
                T.para(
                    tf,
                    foot,
                    size=size - 1,
                    color=T.CRIMSON,
                    italic=True,
                    space_before=8,
                    space_after=0,
                )

        T.fitted(sl, x + 0.18, body_top + 0.14, cw - 0.36, ch - 0.56 - 0.24, write, 13, 8.5)
    top = y + ch + GAP + 0.05
    if s.get("items"):
        _bullets(sl, top, T.BODY_BOTTOM, s["items"], s.get("size", 13))


GRADE = {"Yes": T.GREEN, "Partial": T.AMBER, "No": T.MUTED}


def compare(s: dict[str, Any]) -> None:
    """A capability matrix: Yes, Partial and No graded by colour, one column highlighted."""
    sl, y = T.content(s["title"])
    bottom = _band(sl, s["insight"], T.BODY_BOTTOM, "") if s.get("insight") else T.BODY_BOTTOM
    if s.get("intro"):
        y = _question(sl, y, s["intro"])
    T.fitted_table(
        sl, s["rows"], T.ML, y, T.CW, bottom - y, s.get("col_w"), start=s.get("size", 12)
    )
    tbl = sl.shapes[-1].table
    us = s.get("us")
    for r, row in enumerate(s["rows"]):
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            run = cell.text_frame.paragraphs[0].runs[0]
            if r and c and text in GRADE:
                run.font.color.rgb = GRADE[text]
                run.font.bold = True
            if c:
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            if r and c == us:
                cell.fill.solid()
                cell.fill.fore_color.rgb = T.PINK_L


def shot(s: dict[str, Any]) -> None:
    """A real screenshot from MAYA, framed, with its commentary beside it."""
    from PIL import Image

    sl, y = T.content(s["title"])
    path = SCREENS / s["image"]
    with Image.open(path) as im:
        ratio = im.height / im.width
    bottom = T.BODY_BOTTOM - 0.32
    iw = T.CW * s.get("width", 0.6)
    ih = iw * ratio
    if ih > bottom - y:
        ih = bottom - y
        iw = ih / ratio
    pic = sl.shapes.add_picture(str(path), In(T.ML), In(y), In(iw), In(ih))
    pic.line.color.rgb = T.RULE
    pic.line.width = In(0.01)
    tf = T.txt(sl, T.ML, y + ih + 0.08, iw, 0.26)
    T.para(tf, s["caption"], size=9.5, color=T.MUTED, italic=True, first=True, space_after=0)
    x = T.ML + iw + 0.4
    w = T.ML + T.CW - x

    def write(tf: Any, size: float) -> None:
        for i, it in enumerate(s["items"]):
            head, body = it if isinstance(it, tuple) else (None, it)
            parts = [("▪  ", T.CRIMSON, True)]
            parts += [(head + "  ", T.CRIMSON_D, True)] if head else []
            T.runs(
                tf,
                parts + [(body, T.SLATE, False)],
                size=size,
                first=i == 0,
                space_after=0,
                space_before=0 if i == 0 else 7,
            )

    T.fitted(sl, x, y, w, T.BODY_BOTTOM - y, write, s.get("size", 14), 8.5)


KINDS = {
    "cover": cover,
    "thanks": thanks,
    "tldr": tldr,
    "numbered": numbered,
    "quotes": quotes,
    "arch": arch,
    "lanes": lanes,
    "columns": columns,
    "compare": compare,
    "shot": shot,
}
