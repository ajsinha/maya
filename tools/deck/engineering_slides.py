"""
MAYA — Model & AI Lifecycle Assurance
Copyright © 2026 Ashutosh Sinha <ajsinha@gmail.com>. All rights reserved.
Proprietary and confidential. See LICENSE and NOTICE at the repository root.
"""
# -*- coding: utf-8 -*-
"""MAYA — Model and Feature Engineering deck.

The practitioner's deck: how a feature is defined, how a featureset is composed,
how a warrant carries both to an execution engine, and what comes back. Same
Harvard-Crimson theme as the other two (see theme.py).
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
exec(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "theme.py")).read())

MONO = "Consolas"


def code(sl, x, y, w, lines, fs=9.5, title=None):
    """Monospace block sized from its own content."""
    LINE = 1.34
    lh = fs * LINE * 1.26 / 72.0
    head = 0.32 if title else 0.0
    h = head + len(lines) * lh + 0.30
    rect(sl, x, y, w, h, fill=RGBColor(0xF4, 0xF2, 0xEF))
    rect(sl, x, y, 0.045, h, fill=SLATE)
    if title:
        tf = txt(sl, x + 0.22, y + 0.10, w - 0.4, 0.24)
        para(tf, title, size=9, color=CRIMSON, bold=True, first=True, space_after=0)
    tf = txt(sl, x + 0.22, y + head + 0.11, w - 0.4, h - head - 0.2)
    for i, ln in enumerate(lines):
        col = MUTED if ln.strip().startswith("#") else INK
        para(tf, ln if ln else " ", size=fs, color=col, font=MONO,
             first=(i == 0), space_after=0, line=1.34)
    return h


def steps(sl, x, y, w, items, h=1.45):
    """A numbered row of stages. The deck's way of saying 'in this order'."""
    n = len(items)
    gap = 0.20
    bw = (w - gap * (n - 1)) / n
    for i, (num, title, body) in enumerate(items):
        xx = x + i * (bw + gap)
        rect(sl, xx, y, bw, h, fill=WHITE, line=RULE)
        rect(sl, xx, y, bw, 0.05, fill=CRIMSON)
        tf = txt(sl, xx + 0.16, y + 0.16, bw - 0.32, h - 0.3)
        para(tf, num, size=14, color=CRIMSON, bold=True, font=SERIF,
             first=True, space_after=2)
        para(tf, title, size=11, color=INK, bold=True, space_after=4)
        para(tf, body, size=8.5, color=SLATE, line=1.15)


def arrow(sl, x, y, w, h, fill, line=None):
    """A right-pointing chevron. A process reads left to right, and a shape
    that points is worth more than a line that has to be followed."""
    sh = sl.shapes.add_shape(MSO_SHAPE.CHEVRON, In(x), In(y), In(w), In(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.0)
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.16
    sh.text_frame.word_wrap = True
    sh.text_frame.margin_left = In(0.20)
    sh.text_frame.margin_right = In(0.10)
    sh.text_frame.margin_top = sh.text_frame.margin_bottom = In(0.05)
    return sh


def down(sl, x, y, h, color=CRIMSON, w=0.22):
    """A short downward arrow: one lane handing something to the next."""
    sh = sl.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, In(x), In(y), In(w), In(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def lane(sl, x, y, w, h, label, sub, tint, accent):
    """A swimlane: who is doing this, and the band their work happens in."""
    rect(sl, x, y, w, h, fill=tint)
    rect(sl, x, y, 0.05, h, fill=accent)
    tf = txt(sl, x + 0.16, y + 0.14, 1.02, h - 0.28)
    para(tf, label, size=10.5, color=accent, bold=True, font=SERIF,
         first=True, space_after=2)
    para(tf, sub, size=7.6, color=SLATE, line=1.12, space_after=0)


def stage(sl, x, y, w, h, title, body, fill, edge, text=INK):
    """One step in a lane, drawn as a chevron so the direction is the shape."""
    sh = arrow(sl, x, y, w, h, fill, edge)
    tf = sh.text_frame
    para(tf, title, size=9.5, color=text, bold=True, font=SERIF, first=True,
         space_after=1)
    para(tf, body, size=7.4, color=text if text != INK else SLATE, line=1.10,
         space_after=0)


def listbox(sl, x, y, w, title, items, sub="", accent=CRIMSON, fill=WHITE,
            row=0.235, title_size=11):
    """A named box holding a list of members, each optionally marked.

    ``items`` is a list of (text, colour, marker). The marker is a short label
    printed on the right — 'new', 'gone', 'was UKRPI' — because the interesting
    thing about a composed object is never the list, it is which entries somebody
    here decided.
    """
    h = 0.44 + len(items) * row + (0.20 if sub else 0.0)
    rect(sl, x, y, w, h, fill=fill, line=RULE)
    rect(sl, x, y, w, 0.05, fill=accent)
    tf = txt(sl, x + 0.14, y + 0.13, w - 0.28, 0.26)
    para(tf, title, size=title_size, color=INK, bold=True, font=SERIF,
         first=True, space_after=0)
    yy = y + 0.40
    if sub:
        tf = txt(sl, x + 0.14, yy, w - 0.28, 0.20)
        para(tf, sub, size=7.8, color=SLATE, italic=True, first=True,
             space_after=0)
        yy += 0.20
    for text, colour, marker in items:
        tf = txt(sl, x + 0.14, yy, w - 0.28, row)
        parts = [(text, colour, colour is not SLATE)]
        if marker:
            parts.append((f"   {marker}", CRIMSON, False))
        runs(tf, parts, size=8.8, first=True, space_after=0)
        yy += row
    return h


def note(sl, x, y, w, h, lead, rest, tail=""):
    """A parchment block with a crimson lead-in — the deck's aside."""
    rect(sl, x, y, w, h, fill=PARCH)
    rect(sl, x, y, 0.045, h, fill=CRIMSON)
    tf = txt(sl, x + 0.26, y + 0.13, w - 0.5, h - 0.24)
    parts = [(lead, CRIMSON, True), (rest, INK, False)]
    if tail:
        parts.append((tail, INK, True))
    runs(tf, parts, size=10.5, first=True, space_after=0, line=1.24)


# ============================================================ TITLE
_state["n"] = 0
sl = blank()
rect(sl, 0, 0, SW, SH, fill=WHITE)
rect(sl, 0, 0, SW, 4.35, fill=CRIMSON)
rect(sl, 0, 4.35, SW, 0.06, fill=GOLD)
rect(sl, 0, 0, 0.20, 4.35, fill=CRIMSON_D)
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..",
                    "assets", "logo", "maya-mark-white.png")
if os.path.exists(LOGO):
    sl.shapes.add_picture(LOGO, In(ML + 0.30), In(0.70), In(0.80), In(0.80))
tf = txt(sl, ML + 1.24, 0.99, CW - 1.04, 0.34)
para(tf, "MAYA  ·  MODEL & AI LIFECYCLE ASSURANCE PLATFORM",
     size=11, color=RGBColor(0xE8, 0xB8, 0xC0), bold=True, first=True, space_after=0)
tf = txt(sl, ML + 0.3, 1.62, CW * 0.88, 2.0)
para(tf, "Model and Feature Engineering", size=42, color=WHITE, font=SERIF,
     first=True, space_after=4)
rect(sl, ML + 0.3, 3.28, 1.7, 0.035, fill=RGBColor(0xE8, 0xB8, 0xC0))
tf = txt(sl, ML + 0.3, 3.54, CW * 0.84, 0.8)
para(tf, "How a feature is defined, how a featureset is composed, how a warrant "
         "carries both to an execution engine — and what has to come back",
     size=14, color=RGBColor(0xF4, 0xDF, 0xE3), italic=True, first=True,
     space_after=0, line=1.25)
tf = txt(sl, ML + 0.3, 4.90, CW * 0.55, 1.0)
para(tf, "Ashutosh Sinha", size=20, color=INK, bold=True, font=SERIF,
     first=True, space_after=3)
para(tf, "Independent Researcher", size=12, color=CRIMSON, space_after=1)
para(tf, "September 2026   ·   the practitioner's deck", size=10.5, color=MUTED)
x0 = ML + CW * 0.55
tf = txt(sl, x0, 4.82, CW * 0.45, 1.90)
para(tf, "CHAPTERS", size=9.5, color=CRIMSON, bold=True, first=True, space_after=6)
for i, c in enumerate(["The shape of the whole thing",
                       "Engineering a feature",
                       "Composing a featureset",
                       "Warrants, and what comes back",
                       "What MAYA refuses, and why",
                       "A worked example, with the numbers"], 1):
    runs(tf, [(f"{i}   ", CRIMSON, True), (c, SLATE, False)], size=10.5,
         space_after=3)
footer(sl)


# ============================================================ CH 1
divider("1", "The Shape of the Whole Thing",
        "Features, featuresets, warrants and parameters — and how they meet.",
        ["The process end to end", "What each object is for",
         "Where the pins are", "The two clocks"])

# ------------------------------------------------------------- the five words
sl, y = content("Five words, and how they hold together",
                "The shape of the whole thing")

# --- the equation is the organising device, so it goes first --------------
rect(sl, ML, y, CW, 1.26, fill=PARCH)
rect(sl, ML, y, 0.05, 1.26, fill=CRIMSON)
tf = txt(sl, ML, y + 0.14, CW, 0.62, align=PP_ALIGN.CENTER)
runs(tf, [("f", CRIMSON, True), (" :  ", INK, False),
          ("P", CRIMSON, True), ("  \u2297  ", INK, False),
          ("X", CRIMSON, True), ("  \u2192  ", INK, False),
          ("D(Y)", CRIMSON, True)],
     size=30, first=True, space_after=0)
EQ = [("f", "the KERNEL", "the shape of the computation"),
      ("P", "the PARAMETERS", "the numbers it runs with"),
      ("X", "the FEATURES", "what it reads"),
      ("D(Y)", "the OUTPUT", "a distribution, not a number")]
qw = (CW - 0.5) / 4
for i, (sym, name, gloss) in enumerate(EQ):
    xx = ML + 0.25 + i * qw
    tf = txt(sl, xx, y + 0.82, qw - 0.1, 0.40, align=PP_ALIGN.CENTER)
    runs(tf, [(f"{sym}  ", CRIMSON, True), (name, INK, True)],
         size=9.5, first=True, space_after=1)
    para(tf, gloss, size=8.2, color=SLATE, space_after=0, align=PP_ALIGN.CENTER)

# --- the five terms, each with the confusion it is usually mistaken for ---
CY = y + 1.50
CH = 2.56
cw = (CW - 4 * 0.16) / 5
TERMS = [
    ("Model", "The register entry: what this is for, who owns it, what tier it "
              "sits at.",
     "Not the code, and not the numbers. Those are its versions and its "
     "parameters."),
    ("Kernel", "One immutable version of the model \u2014 the morphism "
               "f : P \u2297 X \u2192 D(Y), pinned by digest.",
     "Not changed by training. Training picks a point in P; the shape stays."),
    ("Feature", "A meaning with an entity, a type, a shape and an owner. A "
                "featureset assembles them into X.",
     "Not a column. A column is where the values happen to sit today."),
    ("Parameters", "An inhabitant of P: the coefficients, weights or constants "
                   "the kernel runs with.",
     "Not a new model. A retrain gives a new parameter set, not a new version."),
    ("Warrant", "A signed, expiring document: this kernel, at that point in P, "
                "over that data, by whom, until when.",
     "Not an execution. MAYA issues it; an engine acts on it and MAYA never "
     "runs anything."),
]
for i, (name, is_, isnt) in enumerate(TERMS):
    xx = ML + i * (cw + 0.16)
    rect(sl, xx, CY, cw, CH, fill=WHITE, line=RULE)
    rect(sl, xx, CY, cw, 0.05, fill=CRIMSON)
    tf = txt(sl, xx + 0.14, CY + 0.16, cw - 0.28, CH - 0.32)
    para(tf, name, size=12, color=INK, bold=True, font=SERIF, first=True,
         space_after=5)
    para(tf, is_, size=8.6, color=INK, line=1.18, space_after=7)
    runs(tf, [("Not  ", CRIMSON, True), (isnt, SLATE, False)],
         size=8.2, line=1.16, space_after=0)

note(sl, ML, CY + CH + 0.20, CW, 0.80,
     "How they correlate. ",
     "A MODEL owns versions; each version is a KERNEL. A FEATURESET presents X. "
     "A fit inhabits P and returns PARAMETERS. A WARRANT is what names a "
     "kernel, a point in P and a source of X together \u2014 ",
     "which is why it is the only one of the five that can be acted on.")

# ---------------------------------------------------------- the big diagram
sl, y = content("Model, feature and warrant management",
                "The shape of the whole thing")

LANE_X = ML
LABEL_W = 1.20
FLOW_X = LANE_X + LABEL_W + 0.10
FLOW_W = CW - LABEL_W - 0.28      # room so the last tip clears the band
LH = 1.06                     # lane height
LGAP = 0.46                   # gap between lanes, where the join arrows sit
STEPS = 5
SGAP = 0.10
BW = (FLOW_W - SGAP * (STEPS - 1)) / STEPS

L1 = y + 0.06
L2 = L1 + LH + LGAP
L3 = L2 + LH + LGAP

TINT1 = RGBColor(0xFA, 0xF4, 0xF5)
TINT2 = RGBColor(0xF6, 0xF3, 0xEF)
TINT3 = RGBColor(0xF3, 0xF4, 0xF6)


def col(i):
    return FLOW_X + i * (BW + SGAP)


# --- lane 1: the data -----------------------------------------------------
lane(sl, LANE_X, L1, CW, LH, "Feature", "what the model reads", TINT1, CRIMSON)
for i, (t, b) in enumerate([
        ("Define", "a meaning: entity, shape, owner"),
        ("Derive", "Z = f(X, Y), lineage recorded"),
        ("Materialise", "values into a pinned Delta version"),
        ("Declare a set", "a SCHEMA of named slots"),
        ("Publish v1", "slots filled, every binding pinned")]):
    stage(sl, col(i), L1 + 0.09, BW, LH - 0.18, t, b,
          WHITE if i < 4 else PARCH_D, RULE)

# --- lane 2: the model ----------------------------------------------------
lane(sl, LANE_X, L2, CW, LH, "Model", "what will run", TINT2, CRIMSON)
for i, (t, b) in enumerate([
        ("Register", "owner, purpose, legal entity"),
        ("Assess", "the tier, which sets the control depth"),
        ("Create version", "the kernel f : P \u00d7 X \u2192 D(Y)"),
        ("Approve", "a quorum whose size follows the tier"),
        ("Point an alias", "refused unless the contract refines")]):
    stage(sl, col(i), L2 + 0.09, BW, LH - 0.18, t, b,
          WHITE if i != 2 else PARCH_D, RULE)

# --- lane 3: execution ----------------------------------------------------
lane(sl, LANE_X, L3, CW, LH, "Execution", "outside MAYA", TINT3, SLATE)
for i, (t, b) in enumerate([
        ("Fit warrant", "model version \u00d7 featureset version \u00d7 window"),
        ("Engine runs", "MAYA signs it and waits"),
        ("Parameter set", "an inhabitant of P comes back"),
        ("Approve it", "by somebody other than who recorded it"),
        ("Score warrant", "names the model AND the parameter set")]):
    stage(sl, col(i), L3 + 0.09, BW, LH - 0.18, t, b,
          WHITE if i not in (0, 4) else PARCH_D, RULE)

# --- the joins: the only vertical arrows, and each one is a claim ---------
JOIN1 = col(4) + BW * 0.42
down(sl, JOIN1, L1 + LH + 0.05, LGAP - 0.10, CRIMSON)
tf = txt(sl, FLOW_X, L1 + LH + 0.11, JOIN1 - FLOW_X - 0.14, 0.26)
para(tf, "the featureset version a fit will read  \u2192", size=8,
     color=CRIMSON, italic=True, first=True, space_after=0,
     align=PP_ALIGN.RIGHT)

JOIN2 = col(2) + BW * 0.42
down(sl, JOIN2, L2 + LH + 0.05, LGAP - 0.10, CRIMSON)
tf = txt(sl, FLOW_X, L2 + LH + 0.11, JOIN2 - FLOW_X - 0.14, 0.26)
para(tf, "the model version it is issued against  \u2192", size=8,
     color=CRIMSON, italic=True, first=True, space_after=0,
     align=PP_ALIGN.RIGHT)

# --- the band everything lands in ----------------------------------------
BAND = L3 + LH + 0.20
rect(sl, ML, BAND, CW, 0.44, fill=CRIMSON)
tf = txt(sl, ML + 0.20, BAND + 0.09, CW - 0.4, 0.30)
runs(tf, [("Evidence chain   ", WHITE, True),
          ("every act above appends a node \u2014 append-only, hash-linked, and "
           "verified by re-deriving each node rather than re-reading it",
           RGBColor(0xF4, 0xDF, 0xE3), False)],
     size=9.5, first=True, space_after=0)

tf = txt(sl, ML, BAND + 0.58, CW, 0.42)
runs(tf, [("Left to right within a lane; downward where one lane hands "
           "something to the next. ", CRIMSON, True),
          ("The two vertical arrows are the whole design: a featureset version "
           "and a model version meet in a fit warrant, and the parameter set "
           "that comes back is what the next warrant names.", SLATE, False)],
     size=9.5, first=True, space_after=0, line=1.2)

# ------------------------------------------------------- what each object is for
sl, y = content("What each object is for", "The shape of the whole thing")
data = [["Object", "Answers", "Pinned by"],
        ["Feature", "what this signal MEANS — entity, shape, owner, definition",
         "nothing; it is a declaration"],
        ["Feature view version", "where the values ARE, at a Delta version that "
         "does not move", "a featureset version"],
        ["Featureset", "what X IS: a schema of named slots a kernel is defined "
         "over", "a model version's contract"],
        ["Featureset version", "which features fill that schema, and from which "
         "exact bytes", "a warrant, a snapshot"],
        ["Model version", "the kernel f : P × X → D(Y) — immutable once approved",
         "a warrant, an alias"],
        ["Warrant", "what may be run, by whom, until when, on what data",
         "signed; nothing pins it"],
        ["Parameter set", "an inhabitant of P: the numbers a fit produced",
         "a scoring warrant"]]
h = table(sl, data, ML, y, CW, col_w=[2.0, 5.7, 3.9], row_h=0.30, fs=10, hfs=10,
          bold_col0=True, first_col_color=CRIMSON)
note(sl, ML, y + h + 0.24, CW, 0.86,
     "Read the third column. ",
     "Every object above is pinned by something further down, and a pin is "
     "always to a version rather than to a name. That is the single structural "
     "decision this platform makes over and over, and it is why ",
     "the same identifier resolves to the same bytes a year later.")

# ------------------------------------------------------------- the two clocks
sl, y = content("Two clocks, never one", "The shape of the whole thing")
h = code(sl, ML, y, CW * 0.56, [
 "event_ts    when the fact was TRUE in the world",
 "ingest_ts   when the platform LEARNED it",
 "",
 "point-in-time rule:",
 "    event_ts  <= label_ts",
 "    ingest_ts <= as_of",
], fs=10)
tf = txt(sl, ML, y + h + 0.26, CW * 0.56, 2.2)
para(tf, "\u201cWhat did we know, and when did we know it\u201d is a different "
         "question from \u201cwhat was true\u201d, and a training set built "
         "without both answers only the second.",
     size=11.5, color=INK, first=True, space_after=8, line=1.26)
para(tf, "A row missing either clock is refused at the upload rather than two "
         "layers later during assembly \u2014 which is where it stops being "
         "fixable.",
     size=11, color=SLATE, space_after=0, line=1.24)

x = ML + CW * 0.60
tf = txt(sl, x, y, CW * 0.40, 0.35)
para(tf, "Where the clocks decide the answer", size=12.5, color=CRIMSON,
     bold=True, font=SERIF, first=True, space_after=0)
data = [["", ""],
        ["Assembly", "the latest fact true by the label date AND known by as_of"],
        ["Normalisation", "statistics fitted only from what was knowable then"],
        ["Derived features", "ingest_ts(Z) = max over the inputs"],
        ["Alignment", "a back-filled value carries the clock it became knowable at"],
        ["Telemetry", "a review of last quarter sees last quarter's population"]]
table(sl, data, x, y + 0.42, CW * 0.40, col_w=[1.35, 3.55], header=False,
      row_h=0.30, fs=9.5, bold_col0=True, first_col_color=CRIMSON)


# ============================================================ CH 2
divider("2", "Engineering a Feature",
        "A meaning, a shape, a lineage, and an owner who answers for it.",
        ["Not always a number", "Derived features", "The leakage rules",
         "Sealing and ownership"])

# ------------------------------------------------------------ dimensionality
sl, y = content("A feature is not always a number", "Engineering a feature")
h = code(sl, ML, y, CW * 0.52, [
 '{"name": "usd_curve", "entity": "book_id",',
 ' "dtype": "numeric",',
 ' "shape": [10],',
 ' "components": ["1m","3m","6m","1y","2y",',
 '                "3y","5y","7y","10y","30y"]}',
], fs=9.5)
data = [["Shape", "Kind", "Example"],
        ["[]", "scalar", "debt service coverage"],
        ["[10]", "vector", "a curve at ten tenors"],
        ["[10, 10]", "matrix", "the correlation between them"],
        ["[5, 10, 10]", "tensor", "that matrix under five scenarios"]]
table(sl, data, ML, y + h + 0.28, CW * 0.52, col_w=[1.3, 1.2, 4.4],
      row_h=0.30, fs=9.5, hfs=9.5, bold_col0=True, first_col_color=CRIMSON)

x = ML + CW * 0.56
tf = txt(sl, x, y, CW * 0.44, 3.4)
para(tf, "The component order IS the axis order", size=12.5, color=CRIMSON,
     bold=True, font=SERIF, first=True, space_after=8)
para(tf, "A curve whose tenors came back alphabetically would be a different "
         "curve \u2014 and one nobody would notice was wrong. Resolution "
         "preserves insertion order and never sorts.",
     size=11, color=INK, space_after=10, line=1.25)
para(tf, "The names are what make composition mean something. \u201cDrop the "
         "50-year point\u201d is an operation on a name; on an anonymous array "
         "it is an operation on an index, and an index is not a meaning.",
     size=11, color=SLATE, space_after=10, line=1.25)
para(tf, "A declared shape is checked against the values that arrive. A shape "
         "nobody verifies is a comment, and a model given ten tenors where it "
         "expected eleven produces an answer rather than an error.",
     size=11, color=SLATE, space_after=0, line=1.25)

# ---------------------------------------------------------- derived features
sl, y = content("Derived features, and the four rules", "Engineering a feature")
h = code(sl, ML, y, CW * 0.50, [
 "lot_to_living_ratio = lot_size_sqft / living_area_sqft",
 "property_age        = year(event_ts) - year_built",
 "log_living_area     = log(living_area_sqft)",
 "",
 "# arithmetic, comparison, nine total functions,",
 "# the row's own clock. whitelisted at the AST,",
 "# so what cannot be expressed cannot be smuggled in.",
], fs=9)
note(sl, ML, y + h + 0.26, CW * 0.50, 1.14,
     "MAYA transforms features it holds; ",
     "it does not run models. Computing x / y over a stored column is the same "
     "class of act as the null rate it already reports for every "
     "materialisation. ",
     "Running a kernel is not.")

x = ML + CW * 0.54
tf = txt(sl, x, y, CW * 0.46, 0.32)
para(tf, "Four rules, each because the alternative fails quietly", size=12,
     color=CRIMSON, bold=True, font=SERIF, first=True, space_after=0)
data = [["Rule", "What it prevents"],
        ["Lineage is the transitive closure",
         "a primitive retired under something deriving from it"],
        ["ingest_ts(Z) = max(inputs)",
         "the easiest way to leak the future: a value appearing knowable "
         "before its inputs were"],
        ["No slot may derive from the label",
         "leakage with a division sign in front of it"],
        ["Certification is the meet",
         "deriving from an uncertified feature to launder it"]]
table(sl, data, x, y + 0.40, CW * 0.46, col_w=[2.3, 3.1], row_h=0.30, fs=9,
      hfs=9, bold_col0=True, first_col_color=CRIMSON)

# ------------------------------------------------------------------ leakage
sl, y = content("The one that gets refused", "Engineering a feature")
h = code(sl, ML, y, CW * 0.54, [
 "feature: price_per_sqft",
 "inputs:  [sale_price, living_area_sqft]",
 "         # sale_price is the featureset's LABEL",
 "",
 "409  'price_per_sqft' is computed from 'sale_price',",
 "     which this featureset declares as its label \u2014 a",
 "     feature derived from the label leaks the answer",
 "     into the training set.",
], fs=9)
tf = txt(sl, ML, y + h + 0.28, CW * 0.54, 1.5)
para(tf, "The check walks the WHOLE lineage, not the immediate inputs: a "
         "derivation of a derivation of the label is still the label.",
     size=11, color=INK, first=True, space_after=8, line=1.25)
para(tf, "And it runs before the view is resolved, so the message is "
         "\u201cyou cannot train on the answer\u201d rather than "
         "\u201cno view supplies that\u201d.",
     size=11, color=SLATE, space_after=0, line=1.25)

x = ML + CW * 0.58
tf = txt(sl, x, y, CW * 0.42, 3.6)
para(tf, "Why this one matters more than it looks", size=12.5, color=CRIMSON,
     bold=True, font=SERIF, first=True, space_after=8)
para(tf, "A model trained on a feature derived from its own label performs "
         "extraordinarily in validation and fails on the first case it has not "
         "already seen.",
     size=11, color=INK, space_after=10, line=1.25)
para(tf, "It is not caught by any metric, because every metric agrees the model "
         "is excellent. It is caught by knowing what the column was made from "
         "\u2014 which is why the lineage is recorded rather than inferred.",
     size=11, color=SLATE, space_after=10, line=1.25)
para(tf, "The same argument produces the ingest-clock rule, the point-in-time "
         "assembly rule, and the refusal to normalise without an as_of. They "
         "are one idea applied four times.",
     size=11, color=SLATE, space_after=0, line=1.25)

# --------------------------------------------------- sealing and ownership
sl, y = content("Sealed, ephemeral, owned", "Engineering a feature")
card(sl, ML, y, CW * 0.315, 2.5, "1", "Sealed",
     "Final: no amendment, no further versions, no change of owner \u2014 and "
     "still composable. That combination is the point. A parent that cannot "
     "move is a parent worth building on, and evolution moves to a child where "
     "it stays visible.")
card(sl, ML + CW * 0.343, y, CW * 0.315, 2.5, "2", "Ephemeral",
     "A time to live, then destroyed. Cannot be sealed \u2014 permanent and "
     "temporary are not two flags that happen to be set. Cannot be composed "
     "from: a child that resolves today and dangles tomorrow. Its rows go; its "
     "evidence stays.")
card(sl, ML + CW * 0.686, y, CW * 0.314, 2.5, "3", "Owned",
     "Two facts kept apart. The CREATOR is history and never moves. The OWNER "
     "is a responsibility, transferred by name with the handover witnessed. An "
     "owner field that quietly becomes a leaver's username is how a model ends "
     "up accountable to nobody.")
note(sl, ML, y + 2.72, CW, 0.92,
     "Sealing is its own permission. ",
     "Whoever may define a thing is not automatically who may end it, so it "
     "sits with the second line. Breaking a seal is administrators-only and "
     "needs a reason \u2014 ",
     "a seal anybody could lift would not be a seal.")


# ============================================================ CH 3
divider("3", "Composing a Featureset",
        "A schema, versions that fill it, and one fold that does both.",
        ["Schema and constituents", "The composition monoid",
         "Rolling forward", "Retrieval policy"])

# ----------------------------------------------------- schema vs constituents
sl, y = content("A featureset declares a schema; a version fills it",
                "Composing a featureset")
h = code(sl, ML, y, CW * 0.56, [
 "featureset: inflation",
 "schema:  gb_index: numeric",
 "         us_index: numeric",
 "         daily_index: numeric",
 "",
 "@v1   gb_index\u2192UKRPI  us_index\u2192USCPI  daily\u2192DAILY_INFL",
 "@v2   gb_index\u2192UKRPI  us_index\u2192USCPI  daily\u2192EUHICP",
], fs=9.5)
tf = txt(sl, ML, y + h + 0.28, CW * 0.56, 1.8)
para(tf, "@v2 draws on a different feature entirely, and a model defined over "
         "inflation does not change: it reads daily_index, and always did.",
     size=11.5, color=INK, first=True, space_after=9, line=1.25)
para(tf, "A version that CANNOT fill the schema is refused. It is not a version "
         "of this featureset \u2014 it is a different one, or it is a model "
         "change, and the register decides which so nobody has to remember.",
     size=11, color=SLATE, space_after=0, line=1.25)

x = ML + CW * 0.60
tf = txt(sl, x, y, CW * 0.40, 0.32)
para(tf, "Every binding pins exactly", size=12.5, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=0)
data = [["", ""],
        ["feature", "which signal fills the slot"],
        ["view", "which view supplies its values"],
        ["view_version", "which version of that view"],
        ["delta_version", "which write to that path"],
        ["namespace", "the resolved Delta location"]]
th = table(sl, data, x, y + 0.40, CW * 0.40, col_w=[1.5, 3.4], header=False,
           row_h=0.29, fs=9.5, bold_col0=True, first_col_color=CRIMSON)
note(sl, x, y + 0.40 + th + 0.18, CW * 0.40, 1.30,
     "A path is mutable. ",
     "A set that named views without pinning them would resolve to different "
     "bytes next month with its digest unchanged \u2014 which is adversarial "
     "finding C-2, ",
     "one level out from the view.")

# --------------------------------------------- worked example: a feature
sl, y = content("Inheriting a feature: a curve with a tenor added",
                "Composing a featureset")

CW3 = (CW - 1.30) / 3
GAPX = 0.65

h1 = listbox(sl, ML, y, CW3, "usd_curve", [
    ("1m", INK, ""), ("3m", INK, ""), ("1y", INK, ""),
    ("5y", INK, ""), ("10y", INK, "")],
    sub="the parent \u2014 sealed, so it cannot move", accent=SLATE,
    fill=PARCH)

x2 = ML + CW3 + GAPX
OPS_H = h1 + 0.24
rect(sl, x2, y, CW3, OPS_H, fill=WHITE, line=RULE)
rect(sl, x2, y, CW3, 0.05, fill=CRIMSON)
tf = txt(sl, x2 + 0.14, y + 0.13, CW3 - 0.28, 0.26)
para(tf, "its own operations", size=11, color=INK, bold=True, font=SERIF,
     first=True, space_after=0)
tf = txt(sl, x2 + 0.14, y + 0.40, CW3 - 0.28, 0.20)
para(tf, "applied last, so they beat everything above", size=7.8, color=SLATE,
     italic=True, first=True, space_after=0)
yy = y + 0.68
for op, target, why in [("add", "30y", "the curve goes further out"),
                        ("drop", "1m", "no longer quoted"),
                        ("override", "3m", "an OIS-based fixing")]:
    tf = txt(sl, x2 + 0.14, yy, CW3 - 0.28, 0.38)
    runs(tf, [(f"{op}  ", CRIMSON, True), (target, INK, True)],
         size=9.5, first=True, space_after=1)
    para(tf, why, size=7.6, color=SLATE, space_after=0)
    yy += 0.40

x3 = x2 + CW3 + GAPX
listbox(sl, x3, y, CW3, "usd_curve_plus", [
    ("3m", INK, "overridden here"), ("1y", SLATE, ""), ("5y", SLATE, ""),
    ("10y", SLATE, ""), ("30y", INK, "added here")],
    sub="1m is gone; the shape follows the components", accent=CRIMSON)

# the arrows between the three columns
down(sl, ML + CW3 + 0.20, y + h1 * 0.42, 0.24, CRIMSON, w=0.24)
down(sl, x2 + CW3 + 0.20, y + h1 * 0.42, 0.24, CRIMSON, w=0.24)

note(sl, ML, y + h1 + 0.28, CW, 0.92,
     "Ask the platform, not the row. ",
     "GET /features/usd_curve_plus/resolved returns the components above AND "
     "where each came from \u2014 3m says it overrode usd_curve. ",
     "The row stores what this feature declares; resolving it is MAYA\u2019s job.")

tf = txt(sl, ML, y + h1 + 1.32, CW, 0.44)
runs(tf, [("The parent is sealed, and that is the point. ", CRIMSON, True),
          ("A sealed feature takes no amendment and can still be composed from, "
           "so a parent that cannot move is a parent worth building on \u2014 "
           "and the change lives in the child, where a reviewer sees it.",
           SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)


# ------------------------------------------ worked example: two featuresets
sl, y = content("Combining two featuresets: the rightmost wins",
                "Composing a featureset")

PW = (CW - 0.60) / 2
hA = listbox(sl, ML, y, PW, "retail_core", [
    ("dscr", INK, ""), ("turnover", INK, ""), ("months_on_book", INK, "")],
    sub="first parent \u2014 least prominent", accent=SLATE, fill=PARCH)
listbox(sl, ML + PW + 0.60, y, PW, "sme_core", [
    ("turnover", INK, "3-year average"), ("sector", INK, ""),
    ("directors", INK, "")],
    sub="second parent \u2014 wins any clash", accent=CRIMSON, fill=PARCH)

MID = y + hA + 0.26
rect(sl, ML, MID, CW, 0.40, fill=CRIMSON)
tf = txt(sl, ML + 0.20, MID + 0.08, CW - 0.4, 0.26)
runs(tf, [("compose: [retail_core, sme_core]", WHITE, True),
          ("     folded left to right, later wins     ", RGBColor(0xF4,0xDF,0xE3), False),
          ("then this set\u2019s own operations, last of all",
           RGBColor(0xF4,0xDF,0xE3), False)],
     size=10, first=True, space_after=0)

RES = MID + 0.62
listbox(sl, ML, RES, PW, "sme_retail", [
    ("dscr", SLATE, "from retail_core"),
    ("turnover", INK, "from sme_core \u2014 it overrode retail\u2019s"),
    ("months_on_book", SLATE, "from retail_core"),
    ("sector", SLATE, "from sme_core"),
    ("directors", SLATE, "from sme_core"),
    ("guarantee_cover", INK, "added by this set")],
    sub="what a kernel defined over it actually reads", accent=CRIMSON)

x = ML + PW + 0.60
tf = txt(sl, x, RES, PW, 2.4)
para(tf, "Only one slot clashed", size=12, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=8)
para(tf, "Both parents declare turnover. sme_core is named second, so its "
         "3-year average is what the set holds \u2014 and the resolved set "
         "SAYS SO, slot by slot.",
     size=10, color=INK, space_after=10, line=1.24)
para(tf, "That is the only rule there is. If a parent could beat a child, "
         "naming a parent would be an act of surrender; if the order did not "
         "matter, \u2018combine these two\u2019 would be ambiguous whenever "
         "they disagreed.",
     size=10, color=SLATE, space_after=10, line=1.24)
runs(tf, [("A child inherits the parents\u2019 retrieval policy the same way",
           CRIMSON, True),
          (" \u2014 one fold, applied to slots, to components and to policy.",
           SLATE, False)],
     size=10, space_after=0, line=1.24)


# ------------------------------------------ modifying something with children
sl, y = content("Modifying something other things are built on",
                "Composing a featureset")

BW4 = (CW - 3 * 0.22) / 4
for i, (num, title, body) in enumerate([
    ("1", "Amend it",
     "Changes the definition and advances its version. Allowed while the "
     "feature is open, and every child records which version it composed "
     "against."),
    ("2", "Seal it",
     "No amendment, no further versions, no change of owner \u2014 and still "
     "composable. A parent that cannot move is a parent worth building on."),
    ("3", "Compose a child",
     "The way to change a sealed thing. The change lives where a reviewer sees "
     "it, beside the parent it departs from."),
    ("4", "Roll forward",
     "For a featureset: mint a version re-resolved to the newest view "
     "versions, with a diff naming every slot that moved."),
]):
    xx = ML + i * (BW4 + 0.22)
    card(sl, xx, y, BW4, 2.15, num, title, body)

DY = y + 2.38
rect(sl, ML, DY, CW, 1.02, fill=PARCH)
rect(sl, ML, DY, 0.05, 1.02, fill=CRIMSON)
tf = txt(sl, ML + 0.26, DY + 0.13, CW - 0.5, 0.80)
para(tf, "And if a parent is amended anyway?", size=11.5, color=CRIMSON,
     bold=True, font=SERIF, first=True, space_after=5)
runs(tf, [("The child says so. ", INK, True),
          ("A composition records the parent\u2019s definition version at the "
           "moment it resolved, so a parent that has moved since shows up as "
           "drift on every child that reads it \u2014 ",
           INK, False),
          ("composed against v1, now at v3.", CRIMSON, True),
          ("  Not a failed read: a child whose parent has moved is something to "
           "be told about, not something that should stop working. But it is "
           "never silent, because a stable name over moving contents is the "
           "failure this platform was built around.", SLATE, False)],
     size=9.8, space_after=0, line=1.22)

tf = txt(sl, ML, DY + 1.20, CW, 0.40)
runs(tf, [("Which is why sealing and composing are the same idea from two "
           "sides. ", CRIMSON, True),
          ("A sealed parent cannot drift, because the amendment that would "
           "move it is refused \u2014 so sealing turns a promise about "
           "stability into a property of the object.", SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)


# ------------------------------------------------------------ the monoid
sl, y = content("Composition: one fold, two objects", "Composing a featureset")
h = code(sl, ML, y, CW * 0.54, [
 "compose([a, b, c])  =  merge(merge(a, b), c)",
 "                       then the object's OWN operations",
 "",
 "merge is associative;  {} is its identity",
 "  \u21d2 composition is a MONOID",
 "  \u21d2 'a combination of features is a feature'",
 "     is a statement, not an aspiration",
], fs=9.5)
tf = txt(sl, ML, y + h + 0.26, CW * 0.54, 1.5)
para(tf, "Left to right, the rightmost wins, and an object\u2019s own "
         "operations are applied last \u2014 because if a parent could "
         "override a child, naming a parent would be an act of surrender.",
     size=11, color=INK, first=True, space_after=8, line=1.25)
para(tf, "Inheriting from one parent and combining several are the same "
         "operation at different arities, so there is one mechanism and "
         "inheritance is the one-parent case.",
     size=11, color=SLATE, space_after=0, line=1.25)

x = ML + CW * 0.58
tf = txt(sl, x, y, CW * 0.42, 0.32)
para(tf, "Every operation is total", size=12.5, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=0)
data = [["Refused", "Because"],
        ["drop of something absent",
         "a drop that quietly does nothing leaves a child differing from what "
         "its author wrote"],
        ["add of something present",
         "say override \u2014 the two read differently to a reviewer and should"],
        ["override of something absent", "say add"],
        ["a cycle", "there is no fixed point to resolve to"],
        ["an ephemeral parent", "it resolves today and dangles tomorrow"]]
table(sl, data, x, y + 0.40, CW * 0.42, col_w=[2.0, 3.5], row_h=0.29, fs=9,
      hfs=9, bold_col0=True, first_col_color=CRIMSON)

# ------------------------------------------------------------- retrieval
sl, y = content("What MAYA does to values on the way out",
                "Composing a featureset")
steps(sl, ML, y, CW, [
    ("1", "Align", "onto a chosen axis, filling gaps by a stated rule"),
    ("2", "Fit", "statistics from what was knowable at the as_of \u2014 on "
                 "OBSERVED values, before anything is filled"),
    ("3", "Fill", "the gaps, reporting how much of each column was invented"),
    ("4", "Normalise", "everything, including what was filled, and return the "
                       "statistics with the data"),
], h=1.42)
tf = txt(sl, ML, y + 1.62, CW * 0.52, 1.6)
para(tf, "The order is the point.", size=12, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=6)
para(tf, "Fitting the normalisation AFTER filling would shrink the spread by "
         "exactly the amount that was invented, because every filled cell sits "
         "at the centre and pulls the variance down.",
     size=11, color=INK, space_after=0, line=1.25)
x = ML + CW * 0.56
tf = txt(sl, x, y + 1.62, CW * 0.44, 1.6)
para(tf, "And the as_of is not optional.", size=12, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=6)
para(tf, "A z-score fitted over the whole column encodes what the mean turned "
         "out to be, including the part that had not happened when the row was "
         "scored. The leaky answer is the one somebody would get by accident, "
         "so it is the one that must not be the default.",
     size=11, color=INK, space_after=0, line=1.25)
note(sl, ML, y + 3.42, CW, 0.92,
     "Attached as default behaviour, overridden by the request. ",
     "Parents, then the object, then the caller \u2014 the same left-to-right "
     "fold, merged column by column, so a parent that fills three columns and a "
     "child that normalises one ",
     "end up doing both.")

# --------------------------------------------------------------- alignment
sl, y = content("Aligning onto an axis, and the leakage that cannot hide",
                "Composing a featureset")
data = [["Rule", "Fills from", "Safe for training"],
        ["flat_forward", "the last observation", "yes"],
        ["flat_backward", "the next observation", "no"],
        ["linear", "both neighbours", "no"],
        ["nearest", "whichever is closer", "no"]]
h = table(sl, data, ML, y, CW * 0.50, col_w=[1.6, 2.9, 2.2], row_h=0.30,
          fs=10, hfs=10, bold_col0=True, first_col_color=CRIMSON)
tf = txt(sl, ML, y + h + 0.26, CW * 0.50, 1.8)
para(tf, "The last three reach into the future. They are right for drawing a "
         "curve or for an explicitly retrospective backtest, and wrong for "
         "training.",
     size=11, color=INK, first=True, space_after=8, line=1.25)
para(tf, "They are not refused. They are stamped honestly.",
     size=11.5, color=CRIMSON, bold=True, space_after=0, line=1.25)

x = ML + CW * 0.54
h2 = code(sl, x, y, CW * 0.46, [
 "grid  t=0   100.0   ingest_ts 0.0",
 "grid  t=1   400.0   ingest_ts 3.0   \u2190 filled from t=3",
 "grid  t=2   400.0   ingest_ts 3.0",
 "grid  t=3   400.0   ingest_ts 3.0",
 "",
 "# a point-in-time read at as_of = 1",
 "# excludes rows 2 and 3 by the ORDINARY rule",
], fs=9)
tf = txt(sl, x, y + h2 + 0.26, CW * 0.46, 1.6)
para(tf, "A value derived from a later observation inherits that "
         "observation\u2019s ingest clock, because that is genuinely when it "
         "became knowable.",
     size=11, color=INK, first=True, space_after=8, line=1.25)
para(tf, "The leakage is not caught by a check. It is made arithmetically "
         "impossible to hide, by the bitemporal machinery already here.",
     size=11, color=SLATE, space_after=0, line=1.25)


# ============================================================ CH 4
divider("4", "Warrants, and What Comes Back",
        "A featureset version and a model version meet; a parameter set returns.",
        ["The fit warrant", "The parameter object", "The scoring warrant",
         "Three routes to P"])

# ---------------------------------------------------------- the fit warrant
sl, y = content("The fit warrant", "Warrants, and what comes back")
h = code(sl, ML, y, CW * 0.54, [
 '"operation":  {"verb": "fit"},',
 '"parameters": {"kind": "estimated_coefficients",',
 '               "source": {"binding": "to_be_fitted"}},',
 '"data": {',
 '  "inputs":  [{"binding":   "featureset",',
 '               "featureset": "nj_home_core", "version": 1,',
 '               "window": {"from": "2019-01-01",',
 '                          "to":   "2024-12-31"},',
 '               "as_of":  "2025-01-15"}],',
 '  "outputs": [{"sink": "parameter_object"}]}',
], fs=9)
note(sl, ML, y + h + 0.24, CW * 0.54, 1.00,
     "The set fixes the columns; the warrant fixes the period. ",
     "That is why window and as_of are here and not in the featureset \u2014 ",
     "one set trains 2019\u201323 and 2020\u201324 without becoming two sets.")

x = ML + CW * 0.58
tf = txt(sl, x, y, CW * 0.42, 0.32)
para(tf, "Five checks before it is signed", size=12.5, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=0)
data = [["Law", "Asks"],
        ["L-W1", "is 'fit' meaningful for this class? T0 has nothing to fit, and "
                 "the refusal says so"],
        ["L-W10", "does the featureset provide what the kernel declares it reads?"],
        ["L-W3", "is the binding one that can answer 'what was known at t'?"],
        ["L-W9", "is the read bounded in BOTH clocks?"],
        ["L-W4", "does it say where the parameters go?"]]
th = table(sl, data, x, y + 0.40, CW * 0.42, col_w=[0.95, 4.55], row_h=0.29, fs=9,
           hfs=9, bold_col0=True, first_col_color=CRIMSON)
tf = txt(sl, x, y + 0.40 + th + 0.20, CW * 0.42, 1.1)
para(tf, "MAYA does not fit anything. It signs this and waits.",
     size=11.5, color=INK, bold=True, first=True, space_after=6, line=1.24)
para(tf, "The estimation happens in an execution engine \u2014 the same "
         "boundary drawn everywhere else here.",
     size=10.5, color=SLATE, space_after=0, line=1.24)

# ------------------------------------------------------- the parameter object
sl, y = content("A fit produces a parameter set, not a model version",
                "Warrants, and what comes back")
h = code(sl, ML, y, CW * 0.52, [
 "f : P \u2297 X \u2192 D(Y)     the kernel",
 "fit                 picks a point in P",
 "",
 "\u21d2 the kernel did not change.",
 "\u21d2 minting a model version per retrain would make",
 "  'the model changed' mean two different things.",
], fs=9.5)
tf = txt(sl, ML, y + h + 0.26, CW * 0.52, 1.9)
para(tf, "But a parameter set does change behaviour \u2014 so it is immutable, "
         "versioned, and an alias cannot point at it until somebody other than "
         "whoever recorded it has approved.",
     size=11, color=INK, first=True, space_after=9, line=1.25)
para(tf, "Accepted ONLY against a warrant MAYA issued, and only when it names "
         "the featureset version that produced it. Without that, \u201cwhich "
         "data produced these numbers\u201d has no answer.",
     size=11, color=SLATE, space_after=0, line=1.25)

x = ML + CW * 0.56
tf = txt(sl, x, y, CW * 0.44, 0.32)
para(tf, "Three routes to P, governed to three depths", size=12.5, color=CRIMSON,
     bold=True, font=SERIF, first=True, space_after=0)
data = [["Route", "Evidence", "Governed"],
        ["fitted", "a warrant MAYA issued", "each set"],
        ["calibrated", "market data, often daily", "the procedure"],
        ["declared", "a person\u2019s assertion", "attestation"]]
table(sl, data, x, y + 0.40, CW * 0.44, col_w=[1.4, 2.5, 1.6], row_h=0.30,
      fs=9.5, hfs=9.5, bold_col0=True, first_col_color=CRIMSON)
tf = txt(sl, x, y + 1.90, CW * 0.44, 2.0)
para(tf, "The middle row is why provenance is not cosmetic: a Hull\u2013White "
         "model recalibrated every morning would drown the register if each day "
         "needed a committee.",
     size=10.5, color=SLATE, first=True, space_after=9, line=1.24)
para(tf, "The bottom row is the model that never trains. A closed form arrives "
         "with its parameters and its kernel is ready from the first day \u2014 "
         "and L-W1 refuses it a fit warrant, so nobody has to remember which "
         "models train.",
     size=10.5, color=SLATE, space_after=0, line=1.24)

# ------------------------------------------------------- the scoring warrant
sl, y = content("The run warrant, and why it is determined",
                "Warrants, and what comes back")
h = code(sl, ML, y, CW * 0.56, [
 '"operation":  {"verb": "score", "deterministic": true},',
 '"parameters": {"kind":   "estimated_coefficients",',
 '               "source": {"binding": "parameter_set",',
 '                          "parameter_set": "01a06d51f9c3",',
 '                          "digest": "sha256:c4e7\u2026"}},',
 '"data": {"inputs":  [{"binding": "request"}],',
 '         "outputs": [{"sink": "response"}]}',
], fs=9)
tf = txt(sl, ML, y + h + 0.28, CW * 0.56, 1.7)
para(tf, "Model version AND parameter set, both pinned. The run is determined: "
         "same kernel, same point in P, same input, same answer.",
     size=11.5, color=INK, first=True, space_after=9, line=1.25)
para(tf, "L-W8 refuses each of to_be_fitted and parameter_set in the "
         "other\u2019s position. A fit that claims to read parameters has the "
         "direction backwards; a score that will not name its inhabitant "
         "produces a number attributable to nothing.",
     size=10.5, color=SLATE, space_after=0, line=1.25)

x = ML + CW * 0.60
tf = txt(sl, x, y, CW * 0.40, 3.6)
para(tf, "What the appraiser gets", size=12.5, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=9)
para(tf, "A number whose provenance runs all the way back:",
     size=11, color=INK, space_after=8, line=1.24)
for step in ["the parameter set that produced it",
             "the featureset version it was fitted from",
             "the pinned Delta namespaces beneath that",
             "the rows that were true and known on a stated date"]:
    runs(tf, [("\u2192  ", CRIMSON, True), (step, SLATE, False)],
         size=10.5, space_after=5)
para(tf, "Not a claim about lineage. A chain of pins, each of which resolves.",
     size=11, color=INK, bold=True, space_before=8, space_after=0, line=1.24)


# ============================================================ CH 5
divider("5", "What MAYA Refuses, and Why",
        "The refusals are the product. Everything else is bookkeeping.",
        ["The refusals that matter", "What it deliberately will not do",
         "Where the boundary is"])

# ------------------------------------------------------------ the refusals
sl, y = content("The refusals that matter", "What MAYA refuses, and why")
data = [["Refused", "Because"],
        ["a training set with either clock unbounded",
         "it cannot be shown point-in-time correct, so it cannot be shown not "
         "to have leaked"],
        ["a feature derived from the label",
         "leakage with a division sign in front of it; every metric will agree "
         "the model is excellent"],
        ["normalising without an as_of",
         "the leaky answer is the one somebody would get by accident, so it "
         "must not be the default"],
        ["a featureset that does not cover the kernel\u2019s inputs",
         "adding a regressor is a model change, not a data change"],
        ["a fit warrant for a T0 model",
         "its parameters come from theory; asking it to fit is a type error"],
        ["parameters with no warrant behind them",
         "\u2018which data produced these numbers\u2019 would have no answer"],
        ["a version approved before its model is assessed",
         "the tier decides how many signatures the approval needs"],
        ["an upload missing entity_id or either clock",
         "accepting it moves the failure two layers from where it was caused"]]
h = table(sl, data, ML, y, CW, col_w=[4.6, 7.0], row_h=0.30, fs=10, hfs=10,
          bold_col0=True, first_col_color=CRIMSON)
note(sl, ML, y + h + 0.22, CW, 0.86,
     "Every refusal carries a remediation. ",
     "\u201cPolicy violation\u201d tells somebody nothing they can act on, so "
     "each of these names what to do instead \u2014 ",
     "a gate people cannot satisfy is a gate people route around.")

# -------------------------------------------------------------- the boundary
sl, y = content("Where the boundary is", "What MAYA refuses, and why")
card(sl, ML, y, CW * 0.315, 2.6, "1", "It does not compute",
     "No fitting, no calibration, no scoring. MAYA issues a warrant and takes "
     "delivery of the result. The captive engine exists so a fresh deployment "
     "can demonstrate the whole governed path against real artifacts \u2014 not "
     "so anybody runs an estate on it.")
card(sl, ML + CW * 0.343, y, CW * 0.315, 2.6, "2", "It does not transform",
     "The expression language is small on purpose. Anything needing a library, "
     "a join or a model is declared external: the definition, the lineage and "
     "the leakage check are still kept, and the register says plainly that the "
     "platform did not compute the values.")
card(sl, ML + CW * 0.686, y, CW * 0.314, 2.6, "3", "It does not choose",
     "No feature selection, no importance ranking, no suggestion engine. It "
     "records what you chose, pins it so it cannot move underneath you, and "
     "refuses the combinations that are type errors.")
note(sl, ML, y + 2.82, CW, 0.94,
     "A platform that claimed to do all three ",
     "would be believed about the parts it does badly. Stating the boundary is "
     "what makes the rest of it worth relying on \u2014 ",
     "and every honest gap is written down in the implementation plan.")

# ============================================================ CH 6
divider("6", "A Worked Example",
        "Real data, real fits: the S&P 500 and the US unemployment rate.",
        ["The data, and its two clocks", "The gap that is really there",
         "Two models over one X", "Extending it, and what that forces"])

# ------------------------------------------------------------- the raw data
sl, y = content("Two real series, on two different clocks",
                "A worked example")

data = [["entity_id", "event_ts", "ingest_ts", "px_close"],
        ["us_equity_index", "2026-08-27", "2026-08-27", "7730.99"],
        ["us_equity_index", "2026-08-28", "2026-08-28", "7711.76"],
        ["us_equity_index", "2026-08-31", "2026-08-31", "7686.14"],
        ["us_equity_index", "2026-09-01", "2026-09-01", "7631.47"],
        ["us_equity_index", "2026-09-02", "2026-09-02", "7666.60"]]
h = table(sl, data, ML, y, CW * 0.46, col_w=[1.85, 1.35, 1.35, 1.2],
          row_h=0.27, fs=8.5, hfs=8.5, bold_col0=True)
tf = txt(sl, ML, y + h + 0.14, CW * 0.46, 0.44)
runs(tf, [("sp500_daily.csv", CRIMSON, True),
          ("   FRED series SP500, 258 rows. The two clocks are equal: an index "
           "close is known the day it happens.", SLATE, False)],
     size=8.8, first=True, space_after=0, line=1.18)

data2 = [["entity_id", "event_ts", "ingest_ts", "u3_rate"],
         ["us_equity_index", "2026-06-30", "2026-07-03", "4.2"],
         ["us_equity_index", "2026-07-31", "2026-08-07", "4.1"],
         ["us_equity_index", "2026-08-31", "2026-09-04", "4.1"]]
h2 = table(sl, data2, ML, y + h + 0.62, CW * 0.46,
           col_w=[1.85, 1.35, 1.35, 1.2], row_h=0.27, fs=8.5, hfs=8.5,
           bold_col0=True)
tf = txt(sl, ML, y + h + h2 + 0.76, CW * 0.46, 0.60)
runs(tf, [("us_unemployment_monthly.csv", CRIMSON, True),
          ("   FRED series UNRATE. ingest_ts is the first Friday of the "
           "following month \u2014 the BLS release convention, computed and "
           "labelled as such rather than passed off as a fetched vintage date.",
           SLATE, False)],
     size=8.8, first=True, space_after=0, line=1.18)

x = ML + CW * 0.50
tf = txt(sl, x, y, CW * 0.50, 0.34)
para(tf, "August\u2019s rate was true on the 31st and known on the 4th",
     size=13, color=CRIMSON, bold=True, font=SERIF, first=True, space_after=8)
para(tf, "A model trained on the 1st of September using August\u2019s figure "
         "is using a number that did not exist yet. It back-tests beautifully "
         "and disappoints in production, and no metric says why.",
     size=11, color=INK, space_after=10, line=1.25)
para(tf, "That is not a hypothetical about leakage. It is the ordinary shape "
         "of every macroeconomic series a bank uses \u2014 published late, "
         "revised afterwards \u2014 and it is why every row here carries two "
         "stamps rather than one.",
     size=11, color=SLATE, space_after=10, line=1.25)
runs(tf, [("Both are scalars, shape []. ", CRIMSON, True),
          ("The curve earlier was a vector; these are not. The shape is "
           "declared either way, because a shape nobody states is a shape "
           "somebody assumes.", SLATE, False)],
     size=11, space_after=0, line=1.25)
note(sl, x, y + 3.18, CW * 0.50, 0.92,
     "The numbers on these slides were not invented. ",
     "Both files were pulled from FRED on 4 September 2026 and ship beside the "
     "deck in docs/examples/ \u2014 ",
     "every figure that follows was computed from them.")


# ---------------------------------------------------- the gap and the fill
sl, y = content("The gap in October is really there", "A worked example")

data = [["month", "S&P 500 close", "return %", "u3 rate", "as the set reads it"],
        ["2025-09", "6688.46", "+3.53", "4.4", "4.4"],
        ["2025-10", "6840.20", "+2.27", "\u2014  not published",
         "4.4   carried forward"],
        ["2025-11", "6849.09", "+0.13", "4.5", "4.5"],
        ["2025-12", "6845.50", "\u22120.05", "4.4", "4.4"],
        ["2026-03", "6528.52", "\u22125.09", "4.3", "4.3"],
        ["2026-04", "7209.01", "+10.42", "4.3", "4.3"],
        ["2026-08", "7686.14", "+2.62", "4.1", "4.1"]]
h = table(sl, data, ML, y, CW * 0.60, col_w=[1.1, 1.55, 1.15, 1.6, 1.58],
          row_h=0.28, fs=9, hfs=9, bold_col0=True, first_col_color=CRIMSON)
tf = txt(sl, ML, y + h + 0.18, CW * 0.60, 0.60)
runs(tf, [("Twelve months, one of them missing. ", CRIMSON, True),
          ("FRED publishes no unemployment rate for October 2025. Nothing was "
           "constructed to make this slide interesting \u2014 the hole is in "
           "the series.", SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)

x = ML + CW * 0.64
h2 = code(sl, x, y, CW * 0.36, [
 "defaults:",
 "  align: {axis: event_ts,",
 "          rule:  flat_forward}",
 "  fill:  {u3_rate: keep}",
 "",
 "# filled: 1 of 12 (8.3%)",
 "# carried from 2025-09,",
 "#   ingest_ts 2025-10-03",
], fs=8.5)
tf = txt(sl, x, y + h2 + 0.22, CW * 0.36, 2.4)
para(tf, "Forward, and only forward", size=12, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=7)
para(tf, "Interpolating October between September and November would use "
         "November\u2019s 4.5 \u2014 published on the 5th of December \u2014 "
         "to fill a row dated the 31st of October.",
     size=10, color=INK, space_after=8, line=1.22)
para(tf, "The platform allows that and stamps it honestly: the filled value "
         "would carry December\u2019s ingest clock, so a point-in-time read in "
         "October excludes it. The leakage is not forbidden, it is made "
         "impossible to hide.",
     size=10, color=SLATE, space_after=8, line=1.22)
runs(tf, [("And the fill rate is part of the answer. ", CRIMSON, True),
          ("One month in twelve is 8%, which is fine. Past a fifth, the report "
           "says so loudly.", SLATE, False)],
     size=10, space_after=0, line=1.22)


# --------------------------------------------------- two models over one X
sl, y = content("Two models, one featureset, both fitted to these twelve months",
                "A worked example")

HW = (CW - 0.40) / 2
h = code(sl, ML, y, HW, [
 "multiple linear regression   \u2014  fitted, n = 10",
 "",
 "next_return  =   43.314",
 "               \u2212  9.708 \u00d7 u3_rate",
 "               \u2212  0.105 \u00d7 prev_return",
 "",
 "kind   estimated_coefficients   \u2192  T2",
 "D(Y)   point_estimate",
 "R\u00b2 = 0.082      residual sd = 4.53 %",
], fs=9)
tf = txt(sl, ML, y + h + 0.22, HW, 1.15)
runs(tf, [("An R\u00b2 of 0.08 is a bad model. ", CRIMSON, True),
          ("Ten monthly observations cannot support three parameters, and the "
           "register stores that verdict beside the coefficients rather than "
           "leaving it in a notebook. A platform that only recorded the "
           "coefficients would make this look like a result.", SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)

x = ML + HW + 0.40
h2 = code(sl, x, y, HW, [
 "GARCH(1,1)   \u2014  maximum likelihood, same returns",
 "",
 "\u03c3\u00b2\u209c  =  3.7684",
 "      + 0.320 \u00d7 \u03b5\u00b2\u209c\u208b\u2081",
 "      + 0.400 \u00d7 \u03c3\u00b2\u209c\u208b\u2081",
 "",
 "kind   estimated_coefficients   \u2192  T2",
 "D(Y)   predictive_distribution",
 "log-likelihood \u221231.90   long-run sd = 3.67 %",
], fs=9)
tf = txt(sl, x, y + h2 + 0.22, HW, 1.15)
runs(tf, [("Persistence \u03b1+\u03b2 = 0.72. ", CRIMSON, True),
          ("Low for a monthly series, which is what twelve observations buy "
           "you. Same features, same entity, same grain \u2014 a different "
           "kernel, a different P, and an output that is a distribution rather "
           "than a number.", SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)

note(sl, ML, y + h + 1.48, CW, 1.10,
     "Two warrants, two parameter sets, one featureset version. ",
     "Each fit warrant names macro_core@v1 and its own model version; each "
     "returns a parameter set recording which featureset version produced it. "
     "Nothing about the data was duplicated to serve two models, and neither "
     "model had to know the other existed \u2014 ",
     "which is what a named presentation of X buys.")


# ---------------------------------------------------- extending the set
sl, y = content("Extending the featureset, and what that forces",
                "A worked example")

h = code(sl, ML, y, CW * 0.47, [
 "featureset: macro_plus",
 "composes:   [macro_core]",
 "operations: [{op: add, name: term_spread}]",
 "",
 "\u2192 macro_plus resolves to FOUR slots",
 "  px_close  u3_rate  prev_return  term_spread",
], fs=9)
tf = txt(sl, ML, y + h + 0.20, CW * 0.47, 1.4)
runs(tf, [("Not a new version of macro_core. ", CRIMSON, True),
          ("Adding a slot changes the schema, and the schema is what a kernel "
           "is defined over \u2014 so it is a new featureset composed from the "
           "old, and macro_core@v1 is untouched for whoever is already using "
           "it.", SLATE, False)],
     size=10, first=True, space_after=0, line=1.22)
note(sl, ML, y + h + 1.66, CW * 0.47, 1.28,
     "The refusal is the feature. ",
     "Widening a featureset under a model never defined over it is the failure "
     "all of this exists to make visible, and it lands at the warrant \u2014 ",
     "before anything is fitted, rather than when somebody notices a "
     "coefficient vector has four entries.")

x = ML + CW * 0.51
h3 = code(sl, x, y, CW * 0.49, [
 "POST /api/v1/fit-warrants",
 "  { model: mlr@1.0.0, featureset: macro_plus }",
 "",
 "409 schema_not_satisfied",
 "    'macro_plus' provides term_spread, which this",
 "    version does not declare in its input schema.",
 "    Bind a featureset whose schema covers the",
 "    kernel's inputs, or create a model version",
 "    whose input schema matches this set \u2014 adding",
 "    a regressor is a MODEL change, not a data one.",
], fs=8.5)
tf = txt(sl, x, y + h3 + 0.20, CW * 0.49, 0.42)
para(tf, "So a third model goes over the wider X:", size=10.5, color=INK,
     bold=True, first=True, space_after=0)
data = [["", ""],
        ["what", "gradient-boosted classifier: drawdown next month"],
        ["kind", "learned_weights  \u2014  T3, not T2"],
        ["fit", "train, not estimate"],
        ["D(Y)", "class_probabilities, not a point estimate"],
        ["reads", "macro_plus@v1, all four slots"]]
table(sl, data, x, y + h3 + 0.66, CW * 0.49, col_w=[0.95, 4.75], header=False,
      row_h=0.27, fs=9, bold_col0=True, first_col_color=CRIMSON)

# ------------------------------------------------------- the data, embedded
sl, y = content("The data itself", "A worked example")

tf = txt(sl, ML, y, CW * 0.52, 2.6)
para(tf, "Both files are inside this deck.", size=13, color=CRIMSON, bold=True,
     font=SERIF, first=True, space_after=9)
para(tf, "Double-click either icon and it opens in Excel. They are the data "
         "every number "
         "in this chapter was computed from, carrying their own provenance in "
         "the header \u2014 which series, from where, retrieved when, and "
         "which column means what.",
     size=11, color=INK, space_after=10, line=1.25)
runs(tf, [("Including what is not certain. ", CRIMSON, True),
          ("The unemployment file says plainly that its ingest_ts is the BLS "
           "release CONVENTION computed from the calendar, not a fetched "
           "vintage date. The lag is real and material; the exact day may be a "
           "day or two out, and a file that did not say so would be inviting "
           "somebody to rely on it.", SLATE, False)],
     size=11, space_after=10, line=1.25)
para(tf, "A deck that quotes figures nobody can check is a deck that has to be "
         "believed. These can be checked.",
     size=11, color=INK, bold=True, space_after=0, line=1.25)

x = ML + CW * 0.56
HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "..", "docs", "examples")

# Embedded as WORKBOOKS, not as OLE packages. An OLE package must be wrapped in
# a compound document, and handing PowerPoint raw bytes under that prog id
# produces an icon that opens nothing -- which is what the first attempt did.
# A .xlsx is a known type and opens on a double-click with no wrapper at all.
import xlsx as _xlsx
from pptx.enum.shapes import PROG_ID

files = [("sp500_daily", "S&P 500, daily close", "S&P 500 daily",
          "FRED series SP500  \u00b7  258 rows  \u00b7  Sep 2025 \u2013 Sep 2026"),
         ("us_unemployment_monthly", "US unemployment rate (U-3)",
          "US unemployment",
          "FRED series UNRATE  \u00b7  12 months  \u00b7  one of them empty")]
yy = y + 0.06
for stem, title, sheet, detail in files:
    source = os.path.join(DATA, stem + ".csv")
    book = os.path.join(DATA, stem + ".xlsx")
    rect(sl, x, yy, CW * 0.44, 1.34, fill=PARCH, line=RULE)
    rect(sl, x, yy, CW * 0.44, 0.05, fill=CRIMSON)
    if os.path.exists(source):
        _xlsx.from_csv(source, book, sheet)
        sl.shapes.add_ole_object(
            book, PROG_ID.XLSX, In(x + 0.20), In(yy + 0.26),
            width=In(0.62), height=In(0.62))
    tf = txt(sl, x + 0.96, yy + 0.22, CW * 0.44 - 1.18, 0.94)
    para(tf, title, size=11.5, color=INK, bold=True, font=SERIF, first=True,
         space_after=3)
    para(tf, detail, size=9, color=SLATE, space_after=3, line=1.18)
    para(tf, stem + ".xlsx", size=8.5, color=CRIMSON, space_after=0)
    yy += 1.52

note(sl, ML, y + 3.22, CW, 0.86,
     "They also ship in the repository, at docs/examples/, as CSV. ",
     "The deck is regenerated from source rather than edited as a binary, so "
     "the figures and the files cannot drift apart \u2014 ",
     "the generator reads these same two CSVs.")


# ============================================================ CLOSING
sl = blank()
rect(sl, 0, 0, SW, SH, fill=CRIMSON)
rect(sl, 0, 0, 0.20, SH, fill=CRIMSON_D)
rect(sl, ML + 0.4, 1.10, 1.7, 0.035, fill=RGBColor(0xE8, 0xB8, 0xC0))
tf = txt(sl, ML + 0.4, 1.40, CW * 0.80, 2.6)
para(tf, "One idea, applied everywhere", size=34, color=WHITE, font=SERIF,
     first=True, space_after=14)
para(tf, "A pin is to a version, never to a name. A clock records when something "
         "became knowable, not when it was written down. A refusal says what to "
         "do instead. And what the platform will not do is written down beside "
         "what it will.",
     size=15, color=RGBColor(0xF4, 0xDF, 0xE3), italic=True, space_after=0,
     line=1.32)
tf = txt(sl, ML + 0.4, 4.22, CW * 0.74, 2.10)
para(tf, "Ashutosh Sinha", size=18, color=WHITE, bold=True, first=True,
     space_after=4)
para(tf, "Independent Researcher   \u00b7   ajsinha@gmail.com", size=12,
     color=RGBColor(0xF2, 0xD8, 0xDC), space_after=14)
para(tf, "Featuresets and parameters: docs/15-featuresets-and-parameters.md   "
         "\u00b7   Composition and retrieval: docs/16-features-composed-and-shaped.md   "
         "\u00b7   Feature platform: docs/07-feature-platform.md",
     size=10.5, color=RGBColor(0xE8, 0xC4, 0xCA), line=1.3)
para(tf, "\u00a9 2026 Ashutosh Sinha. All rights reserved. Proprietary and "
         "confidential \u2014 see LICENSE and NOTICE. Not legal, regulatory or "
         "financial advice.",
     size=8.5, color=RGBColor(0xD8, 0xA0, 0xAC), space_before=10, line=1.25)


import sys
prs.save(sys.argv[1] if len(sys.argv) > 1 else "MAYA-Model-and-Feature-Engineering.pptx")
print("slides:", len(prs.slides._sldIdLst))
