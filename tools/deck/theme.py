"""
The deck design system: Harvard Crimson (#A51C30, spec §16.6), a crimson header
bar with the MAYA mark and a crimson footer bar on every content slide, bold
Calibri titles over a crimson rule, and the layout primitives every slide is
drawn with.

Every primitive that places text measures it with ``metrics`` — the same
estimator the geometry audit uses — and the fitting helpers shrink the type or
refuse, so a slide that would overflow fails the build rather than the reader.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
Proprietary and confidential. See LICENSE and NOTICE at the repository root.
"""

from __future__ import annotations

from typing import Any, Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from pptx.util import Inches as In

from metrics import FOOTER_Y, SAFETY, SANS, SH, SW, est_lines, text_extent, text_h

CRIMSON = RGBColor(0xA5, 0x1C, 0x30)
CRIMSON_D = RGBColor(0x76, 0x14, 0x22)
PINK = RGBColor(0xE8, 0xB8, 0xC0)
PINK_L = RGBColor(0xF6, 0xE6, 0xE9)
INK = RGBColor(0x1C, 0x1C, 0x1E)
SLATE = RGBColor(0x4A, 0x4F, 0x57)
MUTED = RGBColor(0x7A, 0x7F, 0x87)
RULE = RGBColor(0xD8, 0xD4, 0xCF)
PARCH = RGBColor(0xF6, 0xF3, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xB8, 0x8B, 0x3D)
NAVY = RGBColor(0x2E, 0x40, 0x57)

ML = 0.85
CW = SW - 2 * ML
BODY_BOTTOM = FOOTER_Y - 0.12

GREEN = RGBColor(0x1E, 0x7B, 0x4F)
AMBER = RGBColor(0xB0, 0x6A, 0x10)
ROSE = RGBColor(0xC2, 0x3B, 0x52)
WATERMARK = RGBColor(0x8E, 0x1B, 0x2C)  # a shade lighter than CRIMSON_D: the divider's number
HEADER_H = 0.56
TAGLINE = "Model & AI Lifecycle Assurance"
BYLINE = "MAYA  •  Ashutosh Sinha"

_state: dict[str, Any] = {"prs": None, "chapter": "", "n": 0}


class DoesNotFit(Exception):
    """A block that cannot be made to fit its box at the smallest permitted size."""


def new_deck(chapter: str) -> Any:
    """Start a fresh 16:9 presentation and make it the current one."""
    prs = Presentation()
    prs.slide_width = In(SW)
    prs.slide_height = In(SH)
    _state.update(prs=prs, chapter=chapter, n=0)
    return prs


def chapter(name: str) -> None:
    _state["chapter"] = name


def blank() -> Any:
    prs = _state["prs"]
    return prs.slides.add_slide(prs.slide_layouts[6])


def remove(shape: Any) -> None:
    el = shape._element
    el.getparent().remove(el)


def rect(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: Any = None,
    line: Any = None,
    lw: float = 1.0,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def txt(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
) -> Any:
    tb = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(
    tf: Any,
    text: str,
    size: float = 14,
    color: Any = INK,
    bold: bool = False,
    font: str = SANS,
    italic: bool = False,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    space_before: float = 0,
) -> Any:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def runs(
    tf: Any,
    parts: list[tuple],
    size: float = 14,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    level: int = 0,
    space_before: float = 0,
) -> Any:
    """``parts`` is a list of ``(text, color, bold)`` tuples, one run each."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.level = level
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = SANS
    return p


def fitted(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    write: Callable[[Any, float], None],
    start: float,
    floor: float = 9.0,
) -> float:
    """Write a text block at the largest size in [floor, start] whose estimated
    extent fits ``h``; return the height it uses. Raises ``DoesNotFit``."""
    size = start
    while size >= floor:
        tf = txt(sl, x, y, w, h)
        write(tf, size)
        need = text_extent(sl.shapes[-1])
        if need <= h:
            return need
        remove(sl.shapes[-1])
        size -= 0.5
    raise DoesNotFit(f"text block does not fit {w:.2f}x{h:.2f} at {floor}pt")


def mark(sl: Any, x: float, y: float, d: float, color: Any = WHITE) -> None:
    """The MAYA mark, drawn: a diamond inside a ring (python-pptx cannot place the SVG)."""
    ring = rect(sl, x, y, d, d, line=color, lw=1.6, shape=MSO_SHAPE.OVAL)
    ring.fill.background()
    inset = d * 0.24
    rect(
        sl,
        x + inset,
        y + inset,
        d - 2 * inset,
        d - 2 * inset,
        line=color,
        lw=1.6,
        shape=MSO_SHAPE.DIAMOND,
    )


def header(sl: Any, right: str = TAGLINE) -> None:
    """The crimson bar across the top: the mark and the name left, the tagline right."""
    rect(sl, 0, 0, SW, HEADER_H, fill=CRIMSON)
    rect(sl, 0, HEADER_H, SW, 0.03, fill=CRIMSON_D)
    mark(sl, 0.38, 0.11, 0.34)
    tf = txt(sl, 0.86, 0.13, 2.4, 0.3)
    para(tf, "M A Y A", size=14, color=WHITE, bold=True, first=True, space_after=0)
    tf = txt(sl, SW - 4.6, 0.16, 4.2, 0.26, align=PP_ALIGN.RIGHT)
    para(tf, right, size=10.5, color=PINK_L, bold=True, first=True, space_after=0)


def footer(sl: Any) -> None:
    """The crimson bar across the bottom; its gold top edge is the footer rule the audit reads."""
    rect(sl, 0, FOOTER_Y, SW, 0.03, fill=GOLD)
    rect(sl, 0, FOOTER_Y + 0.03, SW, SH - FOOTER_Y - 0.03, fill=CRIMSON)
    tf = txt(sl, ML, FOOTER_Y + 0.17, CW * 0.6, 0.24)
    para(
        tf, "Evidence, not assertion.", size=9, color=PINK_L, italic=True, first=True, space_after=0
    )
    tf = txt(sl, ML + CW * 0.6, FOOTER_Y + 0.17, CW * 0.4, 0.24, align=PP_ALIGN.RIGHT)
    runs(
        tf,
        [(BYLINE + "     ", WHITE, False), (str(_state["n"]), WHITE, True)],
        size=9,
        first=True,
        space_after=0,
    )


def content(title: str, kicker: str | None = None) -> tuple[Any, float]:
    """A content slide's chrome. Returns ``(slide, body_top)``. ``kicker`` is accepted for
    old specs and not drawn: the title carries the slide."""
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=WHITE)
    header(sl)
    footer(sl)
    y = HEADER_H + 0.28
    size = 26.0
    while size > 19 and est_lines(title, CW * SAFETY, size, True, SANS) > 1:
        size -= 1
    th = text_h(title, CW * SAFETY, size, True, SANS, 1.1)
    tf = txt(sl, ML, y, CW, th + 0.04)
    para(tf, title, size=size, color=CRIMSON_D, bold=True, first=True, space_after=0, line=1.1)
    body_top = y + th + 0.30
    rect(sl, ML, body_top - 0.16, CW, 0.022, fill=CRIMSON)
    return sl, body_top


def divider(num: str, title: str, sub: str, points: list[str]) -> Any:
    """A part divider: dark crimson, the part's name on the left under a short bar, and its
    number set huge in the top-right corner a shade lighter than the ground -- a watermark,
    read before the title and never competing with it."""
    _state["chapter"] = f"{num} · {title}"
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=CRIMSON_D)
    rect(sl, 0, 0, SW, 0.09, fill=INK)
    mark(sl, ML, 0.42, 0.34, color=PINK)
    tf = txt(sl, ML + 0.48, 0.45, 2.4, 0.3)
    para(tf, "M A Y A", size=13, color=PINK, bold=True, first=True, space_after=0)
    tf = txt(sl, SW - 6.1, 0.55, 5.6, 4.1, align=PP_ALIGN.RIGHT)
    para(
        tf,
        f"{int(num):02d}" if str(num).isdigit() else str(num),
        size=210,
        color=WATERMARK,
        bold=True,
        first=True,
        space_after=0,
        line=1.0,
    )
    tw = CW * 0.6
    rect(sl, ML, 3.05, 2.4, 0.07, fill=CRIMSON)
    size = 40.0
    while size > 26 and est_lines(title, tw * SAFETY, size, True, SANS) > 2:
        size -= 2
    th = text_h(title, tw * SAFETY, size, True, SANS, 1.05)
    tf = txt(sl, ML, 3.3, tw, th + 0.05)
    para(tf, title, size=size, color=WHITE, bold=True, first=True, space_after=0, line=1.05)

    def write(tf: Any, s: float) -> None:
        para(tf, sub, size=s, color=PINK_L, italic=True, first=True, space_after=12, line=1.25)
        if points:
            para(tf, "   ·   ".join(points), size=s - 2.5, color=PINK, space_after=0, line=1.3)

    top = 3.3 + th + 0.35
    fitted(sl, ML, top, CW * 0.72, FOOTER_Y - 0.25 - top, write, 16, 10)
    return sl


def _row_heights(
    data: list[list[str]], widths: list[float], fs: float, hfs: float, bold_col0: bool
) -> list[float]:
    heights = []
    for r, row in enumerate(data):
        size = hfs if r == 0 else fs
        need = 0.0
        for c, cell in enumerate(row):
            bold = r == 0 or (bold_col0 and c == 0)
            need = max(need, text_h(cell, (widths[c] - 0.18) * SAFETY, size, bold, SANS, 1.0))
        heights.append(max(0.30, need + 0.12))
    return heights


def _fill_cell(
    cell: Any, text: str, r: int, c: int, fs: float, hfs: float, bold_col0: bool
) -> None:
    cell.margin_left = In(0.09)
    cell.margin_right = In(0.07)
    cell.margin_top = cell.margin_bottom = In(0.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = CRIMSON if r == 0 else (PARCH if r % 2 == 0 else WHITE)
    tf = cell.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = str(text)
    run.font.name = SANS
    run.font.size = Pt(hfs if r == 0 else fs)
    run.font.bold = r == 0 or (bold_col0 and c == 0)
    run.font.color.rgb = WHITE if r == 0 else (CRIMSON_D if bold_col0 and c == 0 else INK)


def table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    col_w: list[float] | None = None,
    fs: float = 11.5,
    hfs: float = 11,
    bold_col0: bool = True,
) -> float:
    """A table whose row heights are computed from its wrapped text.
    Returns the rendered height, so callers place what follows from it."""
    cols = len(data[0])
    col_w = col_w or [1.0] * cols
    widths = [w * c / sum(col_w) for c in col_w]
    heights = _row_heights(data, widths, fs, hfs, bold_col0)
    total = sum(heights)
    gf = sl.shapes.add_table(len(data), cols, In(x), In(y), In(w), In(total))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(In(cw)))
    for r, row in enumerate(data):
        tbl.rows[r].height = Emu(int(In(heights[r])))
        for c, cell in enumerate(row):
            _fill_cell(tbl.cell(r, c), cell, r, c, fs, hfs, bold_col0)
    return total


def fitted_table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    col_w: list[float] | None = None,
    start: float = 14.0,
    bold_col0: bool = True,
) -> float:
    """``table`` at the largest size in [9, start] that fits ``h``."""
    fs = start
    while fs >= 9.0:
        used = table(sl, data, x, y, w, col_w, fs=fs, hfs=min(fs, 13), bold_col0=bold_col0)
        if used <= h:
            return used
        remove(sl.shapes[-1])
        fs -= 0.5
    raise DoesNotFit(f"table of {len(data)} rows does not fit {h:.2f}in")


def card(sl: Any, x: float, y: float, w: float, h: float, num: str, title: str, body: str) -> None:
    """A bordered card whose contents are fitted inside the border."""
    rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.9)
    rect(sl, x, y, w, 0.055, fill=CRIMSON)
    inner = w - 0.40
    top = y + 0.18
    if num:
        tf = txt(sl, x + 0.20, top, inner, 0.24)
        para(tf, num, size=10, color=CRIMSON, bold=True, first=True, space_after=0)
        top += 0.28
    used = fitted(
        sl,
        x + 0.20,
        top,
        inner,
        0.72,
        lambda tf, s: para(
            tf,
            title,
            size=s,
            color=CRIMSON_D,
            bold=True,
            font=SANS,
            first=True,
            space_after=0,
            line=1.12,
        ),
        16,
        10.5,
    )
    by = top + used + 0.10
    fitted(
        sl,
        x + 0.20,
        by,
        inner,
        (y + h - 0.14) - by,
        lambda tf, s: para(tf, body, size=s, color=SLATE, first=True, space_after=0, line=1.2),
        14,
        8.5,
    )


def statbar(sl: Any, y: float, stats: list[tuple[str, str]], h: float = 1.25) -> None:
    gap = 0.22
    bw = (CW - gap * (len(stats) - 1)) / len(stats)
    for i, (big, label) in enumerate(stats):
        x = ML + i * (bw + gap)
        rect(sl, x, y, bw, h, fill=PARCH)
        rect(sl, x, y, 0.045, h, fill=CRIMSON)
        fitted(
            sl,
            x + 0.22,
            y + 0.14,
            bw - 0.34,
            0.5,
            lambda tf, s, b=big: para(
                tf,
                b,
                size=s,
                color=CRIMSON,
                bold=True,
                font=SANS,
                first=True,
                space_after=0,
                line=1.0,
            ),
            24,
            14,
        )
        fitted(
            sl,
            x + 0.22,
            y + 0.66,
            bw - 0.34,
            h - 0.74,
            lambda tf, s, lab=label: para(
                tf, lab, size=s, color=SLATE, first=True, space_after=0, line=1.12
            ),
            10.5,
            8,
        )


def connect(
    sl: Any, x1: float, y1: float, x2: float, y2: float, color: Any = SLATE, width: float = 1.5
) -> Any:
    """A straight connector. Elbow connectors auto-route into detours that
    collide with the nodes they join; do not use them."""
    c = sl.shapes.add_connector(1, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c
