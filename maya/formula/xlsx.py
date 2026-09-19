"""
Spreadsheets as first-class models (§29.9): lift an Excel workbook into a formula IR.

MAYA walks the formula graph back from one output cell. Numeric cells it reaches
become the model's inputs, formula cells become lets, the output cell's formula
becomes the body; defined names, where they point at a cell, name it. The v1
scope is deliberately narrow, as the specification asks:

* arithmetic (``+ - * / ^``, unary minus, ``%``) and comparisons;
* SUM, PRODUCT, MIN, MAX, AVERAGE (ranges allowed), ABS, SQRT, EXP, LN, LOG,
  LOG10, POWER, IF, AND, OR, NOT, NORM.S.DIST, NORMSDIST, NORM.DIST;
* named cells and ranges;
* VLOOKUP / HLOOKUP over a table of constants, exact or approximate.

Everything else — text, dates, array formulas, volatile and financial functions,
other lookups, circular references — is **refused, naming the cell and the
construct**, never approximated. An empty cell reads as 0, as in Excel, and says
so in the report.

Then the lift is checked against the workbook itself: every formula cell for
which the file carries a cached result (Excel stores one each time it
recalculates) is compared with the IR evaluated on the workbook's own inputs.
A disagreement blocks submission (``check_formula``); a file that was never
recalculated in Excel is reported as unchecked, not as agreeing.

openpyxl reads the file; with defusedxml installed it refuses XML bombs.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
from dataclasses import dataclass, field
from typing import Any

import numpy as np

from maya.core.errors import CapabilityRefused, ValidationFailed

MAX_BYTES = 20 * 1024 * 1024
MAX_CELLS = 5000
_CELL = r"\$?([A-Za-z]{1,3})\$?(\d+)"
_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<err>\#[A-Z0-9/!?]+[!?]?)
  | (?P<ref>(?:(?P<sheet>'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?[A-Za-z]{1,3}\$?\d+
        (?::\$?[A-Za-z]{1,3}\$?\d+)?)(?![\w(.])
  | (?P<num>(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?)
  | (?P<name>[A-Za-z_\\][\w.]*)
  | (?P<op><>|<=|>=|[-+*/^&=<>%(),;:!{}])
""", re.X)
_CMP = {"=": "eq", "<": "lt", ">": "gt", "<=": "le", ">=": "ge"}
LN10 = math.log(10.0)


def _refuse(where: str, what: str) -> ValidationFailed:
    return ValidationFailed(f"{where}: {what} is outside the spreadsheet import scope and is "
                            "refused rather than approximated", cell=where)


# -- tokens and syntax -------------------------------------------------------------------
def tokenize(text: str, where: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise _refuse(where, f"the text '{text[pos:pos + 12]}'")
        pos = m.end()
        kind = m.lastgroup if m.lastgroup != "sheet" else "ref"
        if m.group("ref"):
            kind = "ref"
        if kind != "ws":
            out.append((kind, m.group(0)))
    return out


class _Parser:
    """Excel precedence, lowest first: comparison, &, + -, * /, ^, %, unary minus."""

    def __init__(self, tokens: list[tuple[str, str]], where: str) -> None:
        self.t, self.i, self.where = tokens, 0, where

    def peek(self) -> tuple[str, str] | None:
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self, value: str | None = None) -> tuple[str, str]:
        tok = self.peek()
        if tok is None or (value is not None and tok[1] != value):
            raise ValidationFailed(f"{self.where}: expected '{value or 'more'}' in the formula")
        self.i += 1
        return tok

    def parse(self) -> Any:
        node = self.compare()
        if self.peek() is not None:
            raise ValidationFailed(f"{self.where}: unexpected '{self.peek()[1]}' in the formula")
        return node

    def _binary(self, ops: tuple[str, ...], nxt: Any) -> Any:
        node = nxt()
        while (tok := self.peek()) and tok[0] == "op" and tok[1] in ops:
            self.take()
            node = ("bin", tok[1], node, nxt())
        return node

    def compare(self) -> Any:
        return self._binary(("=", "<>", "<", ">", "<=", ">="), self.concat)

    def concat(self) -> Any:
        node = self._binary(("&",), self.add)
        return node

    def add(self) -> Any:
        return self._binary(("+", "-"), self.mul)

    def mul(self) -> Any:
        return self._binary(("*", "/"), self.power)

    def power(self) -> Any:
        return self._binary(("^",), self.percent)

    def percent(self) -> Any:
        node = self.unary()
        while (tok := self.peek()) and tok == ("op", "%"):
            self.take()
            node = ("pct", node)
        return node

    def unary(self) -> Any:
        tok = self.peek()
        if tok and tok[0] == "op" and tok[1] in "+-":
            self.take()
            inner = self.unary()
            return ("neg", inner) if tok[1] == "-" else inner
        return self.primary()

    def primary(self) -> Any:
        kind, text = self.take()
        if kind == "num":
            return ("num", float(text))
        if kind == "str":
            return ("str", text[1:-1].replace('""', '"'))
        if kind == "err":
            return ("err", text)
        if kind == "ref":
            return ("ref", text)
        if kind == "name":
            if self.peek() == ("op", "("):
                self.take("(")
                args = []
                if self.peek() != ("op", ")"):
                    args.append(self.compare())
                    while self.peek() in (("op", ","), ("op", ";")):
                        self.take()
                        args.append(self.compare())
                self.take(")")
                return ("call", text.upper().removeprefix("_XLFN."), args)
            return ("name", text)
        if (kind, text) == ("op", "("):
            node = self.compare()
            self.take(")")
            return node
        raise _refuse(self.where, f"'{text}'")


# -- the workbook -------------------------------------------------------------------------
def _col(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _letters(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _cells_of(ref: str) -> tuple[str | None, list[list[str]]]:
    """``Sheet!A1:B3`` -> (sheet, rows of coordinates)."""
    sheet = None
    if "!" in ref:
        sheet, ref = ref.rsplit("!", 1)
        sheet = sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet
    parts = [re.fullmatch(_CELL, p) for p in ref.split(":")]
    (c1, r1), (c2, r2) = (parts[0].groups(), parts[-1].groups())
    cols = range(min(_col(c1), _col(c2)), max(_col(c1), _col(c2)) + 1)
    rows = range(min(int(r1), int(r2)), max(int(r1), int(r2)) + 1)
    return sheet, [[f"{_letters(c)}{r}" for c in cols] for r in rows]


@dataclass
class _Lift:
    formulas: Any
    values: Any
    names: dict[str, tuple[str | None, str]]
    lets: dict[str, Any] = field(default_factory=dict)
    inputs: dict[str, float] = field(default_factory=dict)
    lookups: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    visiting: list[str] = field(default_factory=list)

    # cell access
    def key(self, sheet: str | None, coord: str, here: str) -> str:
        sheet = sheet or here.split("!")[0]
        if sheet not in self.formulas.sheetnames:
            raise ValidationFailed(f"{here}: no sheet named '{sheet}'")
        return f"{sheet}!{coord.replace('$', '').upper()}"

    def raw(self, key: str) -> Any:
        sheet, coord = key.split("!")
        return self.formulas[sheet][coord].value

    def is_formula(self, key: str) -> bool:
        v = self.raw(key)
        return isinstance(v, str) and v.startswith("=")

    def constant(self, key: str, here: str) -> float | None:
        """A non-formula cell as a number; None when empty."""
        v = self.raw(key)
        if v is None:
            return None
        if isinstance(v, bool):
            return float(v)
        if isinstance(v, (int, float)):
            return float(v)
        kind = type(v).__name__ if not isinstance(v, str) else "text"
        raise _refuse(here, f"cell {key} holding {kind} ('{str(v)[:20]}')")

    # the graph walk
    def cell(self, key: str, here: str) -> Any:
        """The IR node standing for one referenced cell."""
        if self.is_formula(key):
            if key in self.visiting:
                raise ValidationFailed("circular reference: " + " -> ".join(
                    self.visiting[self.visiting.index(key):] + [key]))
            if key not in self.lets:
                if len(self.lets) >= MAX_CELLS:
                    raise ValidationFailed(f"more than {MAX_CELLS} formula cells; refused")
                self.visiting.append(key)
                self.lets[key] = self.formula(key)
                self.visiting.pop()
            return {"ref": key}
        value = self.constant(key, here)
        if value is None:
            self.warnings.append(f"{here}: empty cell {key} reads as 0, as in Excel")
            return {"const": 0.0}
        self.inputs[key] = value
        return {"ref": key}

    def formula(self, key: str) -> Any:
        text = self.raw(key)[1:]
        return self.node(_Parser(tokenize(text, key), key).parse(), key)

    def resolve_name(self, name: str, here: str) -> tuple[str | None, str]:
        target = self.names.get(name.upper())
        if target is None:
            raise ValidationFailed(f"{here}: no defined name '{name}' (named cells and ranges "
                                   "are supported; names of formulas or constants are not)")
        return target

    def cells(self, ast: Any, here: str) -> list[str] | None:
        """The cell keys a ref/range/name argument covers, or None for an expression."""
        if ast[0] == "ref":
            sheet, rows = _cells_of(ast[1])
        elif ast[0] == "name" and ast[1].upper() not in ("TRUE", "FALSE"):
            sheet, target = self.resolve_name(ast[1], here)
            _, rows = _cells_of(target)
        else:
            return None
        return [self.key(sheet, c, here) for row in rows for c in row]

    def node(self, ast: Any, here: str) -> Any:
        kind = ast[0]
        if kind == "num":
            return {"const": ast[1]}
        if kind in ("str", "err"):
            raise _refuse(here, f"the {'text' if kind == 'str' else 'error'} literal {ast[1]!r}")
        if kind == "name" and ast[1].upper() in ("TRUE", "FALSE"):
            return {"const": 1.0 if ast[1].upper() == "TRUE" else 0.0}
        if kind in ("ref", "name"):
            keys = self.cells(ast, here)
            if len(keys) != 1:
                raise _refuse(here, f"the range {ast[1]} outside an aggregate or lookup")
            return self.cell(keys[0], here)
        if kind == "neg":
            return {"op": "neg", "args": [self.node(ast[1], here)]}
        if kind == "pct":
            return {"op": "div", "args": [self.node(ast[1], here), {"const": 100.0}]}
        if kind == "bin":
            return self.binary(ast[1], self.node(ast[2], here), self.node(ast[3], here), here)
        if kind == "call":
            return self.call(ast[1], ast[2], here)
        raise _refuse(here, f"'{ast}'")

    @staticmethod
    def binary(op: str, a: Any, b: Any, here: str) -> Any:
        if op == "&":
            raise _refuse(here, "text concatenation (&)")
        if op == "<>":
            return {"op": "where", "args": [{"op": "eq", "args": [a, b]},
                                            {"const": 0.0}, {"const": 1.0}]}
        if op in _CMP:
            return {"op": _CMP[op], "args": [a, b]}
        return {"op": {"+": "add", "-": "sub", "*": "mul", "/": "div", "^": "pow"}[op],
                "args": [a, b]}

    def items(self, args: list[Any], here: str, *, skip_empty: bool) -> list[Any]:
        """Aggregate arguments: ranges expand to their cells (Excel skips empty ones)."""
        out = []
        for a in args:
            keys = self.cells(a, here)
            if keys is None:
                out.append(self.node(a, here))
                continue
            for k in keys:
                if skip_empty and len(keys) > 1 and not self.is_formula(k) and self.raw(k) is None:
                    continue
                out.append(self.cell(k, here))
        if not out:
            raise ValidationFailed(f"{here}: an aggregate over no numbers")
        return out

    def call(self, fn: str, args: list[Any], here: str) -> Any:
        family = next((f for names, f in (
            (("SUM", "PRODUCT", "MIN", "MAX", "AVERAGE"), self._aggregate),
            (("ABS", "SQRT", "EXP", "LN", "LOG10", "LOG", "POWER"), self._maths),
            (("IF", "AND", "OR", "NOT", "TRUE", "FALSE"), self._logic),
            (("NORM.S.DIST", "NORMSDIST", "NORM.DIST"), self._normal),
            (("VLOOKUP", "HLOOKUP"), self.lookup)) if fn in names), None)
        if family is None:
            raise _refuse(here, f"the function {fn}")
        return family(fn, args, here)

    @staticmethod
    def _arity(fn: str, args: list[Any], here: str, lo: int, hi: int) -> None:
        if not lo <= len(args) <= hi:
            raise ValidationFailed(f"{here}: {fn} takes {lo}"
                                   f"{'' if lo == hi else f' to {hi}'} argument(s)")

    def _aggregate(self, fn: str, args: list[Any], here: str) -> Any:
        xs = self.items(args, here, skip_empty=True)
        if fn == "AVERAGE":
            total = xs[0] if len(xs) == 1 else {"op": "add", "args": xs}
            return {"op": "div", "args": [total, {"const": float(len(xs))}]}
        op = {"SUM": "add", "PRODUCT": "mul", "MIN": "min", "MAX": "max"}[fn]
        return xs[0] if len(xs) == 1 else {"op": op, "args": xs}

    def _maths(self, fn: str, args: list[Any], here: str) -> Any:
        if fn == "POWER":
            self._arity(fn, args, here, 2, 2)
            return {"op": "pow", "args": [self.node(args[0], here), self.node(args[1], here)]}
        if fn in ("LOG10", "LOG"):
            self._arity(fn, args, here, 1, 1 if fn == "LOG10" else 2)
            base = {"const": LN10} if len(args) == 1 else \
                {"op": "log", "args": [self.node(args[1], here)]}
            return {"op": "div", "args": [{"op": "log", "args": [self.node(args[0], here)]}, base]}
        self._arity(fn, args, here, 1, 1)
        return {"op": {"LN": "log"}.get(fn, fn.lower()), "args": [self.node(args[0], here)]}

    def _logic(self, fn: str, args: list[Any], here: str) -> Any:
        if fn in ("TRUE", "FALSE"):
            self._arity(fn, args, here, 0, 0)
            return {"const": 1.0 if fn == "TRUE" else 0.0}
        if fn == "IF":
            self._arity(fn, args, here, 2, 3)
            other = self.node(args[2], here) if len(args) == 3 else {"const": 0.0}
            return {"op": "where", "args": [self.node(args[0], here), self.node(args[1], here),
                                            other]}
        if fn == "NOT":
            self._arity(fn, args, here, 1, 1)
            return {"op": "eq", "args": [self.node(args[0], here), {"const": 0.0}]}
        xs = self.items(args, here, skip_empty=True)
        return xs[0] if len(xs) == 1 else {"op": fn.lower(), "args": xs}

    def _normal(self, fn: str, args: list[Any], here: str) -> Any:
        if fn == "NORM.DIST":
            self._arity(fn, args, here, 4, 4)
            x, mu, sd = (self.node(a, here) for a in args[:3])
            z = {"op": "div", "args": [{"op": "sub", "args": [x, mu]}, sd]}
            if self.literal_bool(args[3], fn, here):
                return {"op": "ncdf", "args": [z]}
            return {"op": "div", "args": [{"op": "npdf", "args": [z]}, sd]}
        self._arity(fn, args, here, *((1, 1) if fn == "NORMSDIST" else (2, 2)))
        cumulative = fn == "NORMSDIST" or self.literal_bool(args[1], fn, here)
        return {"op": "ncdf" if cumulative else "npdf", "args": [self.node(args[0], here)]}

    @staticmethod
    def literal_bool(ast: Any, fn: str, here: str) -> bool:
        if ast[0] == "name" and ast[1].upper() in ("TRUE", "FALSE"):
            return ast[1].upper() == "TRUE"
        if ast[0] == "call" and ast[1] in ("TRUE", "FALSE") and not ast[2]:
            return ast[1] == "TRUE"      # TRUE() / FALSE(): how LibreOffice writes the literals
        if ast[0] == "num" and ast[1] in (0.0, 1.0):
            return bool(ast[1])
        raise _refuse(here, f"a {fn} whose TRUE/FALSE argument is not written literally")

    def lookup(self, fn: str, args: list[Any], here: str) -> Any:
        """VLOOKUP/HLOOKUP over a constant table: nested where() on the lookup value."""
        if len(args) not in (3, 4):
            raise ValidationFailed(f"{here}: {fn} takes 3 or 4 arguments")
        table = args[1]
        if table[0] not in ("ref", "name"):
            raise _refuse(here, f"a {fn} table that is not a cell range")
        if table[0] == "ref":
            sheet, rows = _cells_of(table[1])
        else:
            sheet, target = self.resolve_name(table[1], here)
            rows = _cells_of(target)[1]
        if fn == "HLOOKUP":
            rows = [list(col) for col in zip(*rows)]
        if args[2][0] != "num" or args[2][1] != int(args[2][1]):
            raise _refuse(here, f"a {fn} column index that is not a literal whole number")
        idx = int(args[2][1])
        if not 1 <= idx <= len(rows[0]):
            raise ValidationFailed(f"{here}: {fn} index {idx} is outside the table")
        approx = True if len(args) == 3 else self.literal_bool(args[3], fn, here)
        pairs = []
        for row in rows:
            kkey, vkey = (self.key(sheet, row[0], here), self.key(sheet, row[idx - 1], here))
            if self.is_formula(kkey) or self.is_formula(vkey):
                raise _refuse(here, f"a {fn} table containing formulas ({kkey} or {vkey})")
            k, v = self.constant(kkey, here), self.constant(vkey, here)
            if k is not None:
                pairs.append((k, 0.0 if v is None else v))
        if approx and any(a[0] >= b[0] for a, b in zip(pairs, pairs[1:])):
            raise ValidationFailed(f"{here}: an approximate {fn} needs its first "
                                   f"{'column' if fn == 'VLOOKUP' else 'row'} strictly ascending")
        key = self.node(args[0], here)
        node: Any = {"op": "na", "args": []}
        for k, v in (pairs if approx else reversed(pairs)):
            test = {"op": "ge" if approx else "eq", "args": [key, {"const": k}]}
            node = {"op": "where", "args": [test, {"const": v}, node]}
        self.lookups.append({"cell": here, "function": fn, "rows": len(pairs),
                             "match": "approximate" if approx else "exact"})
        return node


# -- naming, output and the check -----------------------------------------------------------
def _defined_names(wb: Any) -> dict[str, tuple[str | None, str]]:
    out: dict[str, tuple[str | None, str]] = {}
    scopes = [(None, wb.defined_names)] + [(ws.title, ws.defined_names) for ws in wb.worksheets]
    for _, names in scopes:
        for name, dn in names.items():
            dests = list(dn.destinations)
            if len(dests) == 1:
                sheet, coord = dests[0]
                out[name.upper()] = (sheet, f"'{sheet}'!{coord.replace('$', '')}")
    return out


def _ident(text: str) -> str:
    s = re.sub(r"\W", "_", text).strip("_") or "cell"
    return f"c_{s}" if s[0].isdigit() else s


def _names_for(keys: list[str], names: dict[str, tuple[str | None, str]]) -> dict[str, str]:
    by_cell: dict[str, str] = {}
    for name, (sheet, target) in names.items():
        s, rows = _cells_of(target)
        if len(rows) == 1 and len(rows[0]) == 1:
            by_cell.setdefault(f"{s}!{rows[0][0]}", name.lower())
    one_sheet = len({k.split("!")[0] for k in keys}) == 1
    out: dict[str, str] = {}
    used: set[str] = set()
    for k in keys:
        base = _ident(by_cell.get(k) or (k.split("!")[1] if one_sheet else k.replace("!", "_")))
        name, n = base, 2
        while name in used:
            name, n = f"{base}_{n}", n + 1
        used.add(name)
        out[k] = name
    return out


def _rename(node: Any, names: dict[str, str]) -> Any:
    if "ref" in node:
        return {"ref": names[node["ref"]]}
    if "op" in node:
        return {"op": node["op"], "args": [_rename(a, names) for a in node["args"]]}
    return node


def _sinks(wb: Any, names: dict[str, tuple[str | None, str]]) -> list[str]:
    """Formula cells no other formula refers to, directly or by name: output candidates."""
    formulas, used = [], set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    key = f"{ws.title}!{c.coordinate}"
                    formulas.append(key)
                    try:
                        toks = tokenize(c.value[1:], key)
                    except ValidationFailed:
                        continue
                    for kind, text in toks:
                        if kind == "name" and text.upper() in names:
                            text = names[text.upper()][1]
                        elif kind != "ref":
                            continue
                        sheet, rows = _cells_of(text)
                        used.update(f"{sheet or ws.title}!{x}" for r in rows for x in r)
    return [f for f in formulas if f not in used]


def _output_key(wb: Any, lift: _Lift, output: str | None) -> str:
    if output:
        here = f"{wb.worksheets[0].title}!A1"
        if lift.names.get(output.upper()):
            keys = lift.cells(("name", output), here)
        else:
            keys = lift.cells(("ref", output), here) if re.fullmatch(
                r"(?:.+!)?\$?[A-Za-z]{1,3}\$?\d+", output) else None
        if not keys or len(keys) != 1:
            raise ValidationFailed(f"'{output}' is not a single cell or a named cell")
        return keys[0]
    sinks = _sinks(wb, lift.names)
    if len(sinks) != 1:
        raise ValidationFailed("Name the output cell: the workbook has "
                               f"{len(sinks) or 'no'} formula cell(s) nothing else uses"
                               + (f" ({', '.join(sinks[:12])})" if sinks else ""),
                               candidates=sinks[:50])
    return sinks[0]


def _check(ir: dict[str, Any], lift: _Lift, cells: dict[str, str], out_key: str
           ) -> dict[str, Any]:
    """Compare the lifted IR, cell by cell, with the results the workbook itself cached."""
    from maya.formula.evaluate import eval_node
    from maya.formula.ir import let_order
    env: dict[str, Any] = {cells[k]: np.array([v]) for k, v in lift.inputs.items()}
    with np.errstate(all="ignore"):
        for name in let_order(ir["lets"]):
            env[name] = eval_node(ir["lets"][name], env, {})
        env[ir["outputs"][0]["name"]] = eval_node(ir["body"], env, {})
    checked, disagreements = 0, []
    for key in list(lift.lets) + [out_key]:
        sheet, coord = key.split("!")
        cached = lift.values[sheet][coord].value
        if cached is None:
            continue
        mine = float(np.asarray(env[cells[key]], dtype=float).reshape(-1)[0])
        if isinstance(cached, str) and cached.startswith("#"):
            same = math.isnan(mine)
        elif isinstance(cached, (int, float)):
            same = bool(np.isclose(mine, float(cached), rtol=1e-9, atol=1e-12))
        else:
            continue
        checked += 1
        if not same:
            disagreements.append({"cell": key, "workbook": cached, "lifted": mine})
    total = len(lift.lets) + 1
    status = "unchecked" if not checked else ("disagreed" if disagreements else "agreed")
    statement = {
        "unchecked": "The workbook carries no cached results (it was never recalculated in "
                     "Excel), so the lift could not be compared with it.",
        "agreed": f"The lifted formula reproduces the workbook's own results on all {checked} "
                  f"of {total} formula cells it caches, at the workbook's inputs.",
        "disagreed": f"{len(disagreements)} of {checked} cached cells disagree with the lifted "
                     "formula; submission is blocked until they agree.",
    }[status]
    return {"status": status, "checked": checked, "formula_cells": total,
            "disagreements": disagreements[:20], "statement": statement}


def lift_workbook(data: bytes, *, output: str | None = None, roles: dict[str, str] | None = None,
                  filename: str = "workbook.xlsx") -> dict[str, Any]:
    """Lift one output cell of an .xlsx into a formula IR, with a report of what was done."""
    from maya.core.backends import has_module
    if not has_module("openpyxl"):
        raise CapabilityRefused("Spreadsheet import needs openpyxl (pip install openpyxl)")
    if len(data) > MAX_BYTES:
        raise ValidationFailed(f"The workbook is over {MAX_BYTES // 2**20} MB; refused")
    import openpyxl
    try:
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is the same answer to the user
        raise ValidationFailed(f"Not a readable .xlsx workbook: {exc}") from exc
    lift = _Lift(formulas, values, _defined_names(formulas))
    out_key = _output_key(formulas, lift, output)
    if not lift.is_formula(out_key):
        raise ValidationFailed(f"The output {out_key} is a value, not a formula")
    lift.visiting.append(out_key)
    body = lift.formula(out_key)
    lift.visiting.pop()
    cells = _names_for(sorted(lift.inputs) + list(lift.lets) + [out_key], lift.names)
    roles = roles or {}
    unknown = set(roles) - {cells[k] for k in lift.inputs}
    if unknown:
        raise ValidationFailed("Roles name inputs the lift does not have: "
                               + ", ".join(sorted(unknown)),
                               inputs=sorted(cells[k] for k in lift.inputs))
    inputs = []
    for key in sorted(lift.inputs, key=lambda k: cells[k]):
        role = roles.get(cells[key], "feature")
        entry: dict[str, Any] = {"name": cells[key], "type": "float64", "role": role}
        if role == "constant":
            entry["value"] = lift.inputs[key]
        inputs.append(entry)
    ir: dict[str, Any] = {
        "outputs": [{"name": cells[out_key], "type": "float64"}], "inputs": inputs,
        "lets": {cells[k]: _rename(v, cells) for k, v in lift.lets.items()},
        "body": _rename(body, cells)}
    from maya.formula.ir import validate_ir
    errors = validate_ir(ir)
    if errors:
        raise ValidationFailed("The lifted IR is invalid", errors=errors)
    check = _check(ir, lift, cells, out_key)
    ir["lifted_from"] = {"workbook": {
        "filename": filename, "sha256": hashlib.sha256(data).hexdigest(), "output": out_key,
        "cells": {cells[k]: k for k in cells},
        "values": {cells[k]: v for k, v in lift.inputs.items()},
        "lookups": lift.lookups, "warnings": sorted(set(lift.warnings)), "check": check}}
    from maya.formula.latex import to_latex
    ir["latex"] = to_latex(ir)
    return ir
